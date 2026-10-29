from datetime import datetime

from openpyxl import load_workbook
from PySide6.QtCore import QCoreApplication, QUrl

from edutrack_desktop.infrastructure.excel_export import export_register
from edutrack_desktop.presentation.backend import DesktopBackend


def test_export_register_groups_courses_by_type_in_separate_sheets(tmp_path):
    manifest = {
        "courses": [
            {
                "id": 1,
                "name": "Pastelería",
                "schedule": "Mañana",
                "group": "1",
                "start_date": "2026-01-15",
                "enrollment_fee": "100.00",
                "monthly_fee": "80.00",
                "capacity": 2,
                "active": True,
            },
            {
                "id": 2,
                "name": "PASTELERÍA",
                "schedule": "Tarde",
                "group": "2",
                "start_date": "2026-01-17",
                "enrollment_fee": "110.00",
                "monthly_fee": "85.00",
                "capacity": 1,
                "active": True,
            },
            {
                "id": 3,
                "name": "Costura",
                "schedule": "Noche",
                "group": "1",
                "start_date": "2026-03-15",
                "enrollment_fee": "90.00",
                "monthly_fee": "70.00",
                "capacity": 5,
                "active": False,
            },
            {
                "id": 4,
                "name": "Bordado",
                "schedule": "Mañana",
                "group": "1",
                "start_date": "2026-02-15",
                "enrollment_fee": "95.00",
                "monthly_fee": "75.00",
                "capacity": 8,
                "active": True,
            },
        ],
        "students": [
            {
                "id": "student-1",
                "last_name": "PÉREZ",
                "first_name": "ANA",
                "document_number": "00123456",
                "phone_numbers": ["987654321"],
                "active": True,
            },
            {
                "id": "student-2",
                "last_name": "GARCÍA",
                "first_name": "LUIS",
                "document_number": "87654321",
                "phone_numbers": ["912345678"],
                "active": True,
            },
            {
                "id": "student-3",
                "last_name": "ALUMNA",
                "first_name": "INACTIVA",
                "document_number": "00000003",
                "phone_numbers": ["900000003"],
                "active": False,
            },
        ],
        "enrollments": [
            {
                "id": "enrollment-1",
                "student_id": "student-1",
                "course_id": 1,
                "enrollment_date": "2026-01-20",
                "enrollment_fee": "100.00",
                "monthly_fee": "80.00",
                "record_code": 101,
                "promoter_ids": [1],
                "active": True,
            },
            {
                "id": "enrollment-2",
                "student_id": "student-2",
                "course_id": 2,
                "enrollment_date": "2026-02-20",
                "enrollment_fee": "110.00",
                "monthly_fee": "85.00",
                "record_code": None,
                "promoter_ids": [],
                "active": True,
            },
            {
                "id": "enrollment-3",
                "student_id": "student-3",
                "course_id": 1,
                "enrollment_date": "2026-01-19",
                "enrollment_fee": "100.00",
                "monthly_fee": "80.00",
                "record_code": 103,
                "promoter_ids": [],
                "active": True,
            },
        ],
        "promoters": [{"id": 1, "name": "PROMOTOR LOCAL"}],
    }
    output = tmp_path / "registro.xlsx"

    export_register(manifest, output)

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["Cursos", "BORDADO", "PASTELERÍA"]
    courses_sheet = workbook["Cursos"]
    assert courses_sheet.max_row == 8
    assert courses_sheet["A1"].value == (
        "SEMANA 1: del 12 al 18 de enero de 2026"
    )
    assert courses_sheet["A2"].value == "Tipo de curso"
    assert courses_sheet["D3"].value == datetime(2026, 1, 15)
    assert courses_sheet["D4"].value == datetime(2026, 1, 17)
    assert courses_sheet["A6"].value == (
        "SEMANA 2: del 9 al 15 de febrero de 2026"
    )
    assert courses_sheet["D8"].value == datetime(2026, 2, 15)
    assert courses_sheet.max_column == 9
    assert courses_sheet["I2"].value == "Cupos disponibles"
    assert {table.ref for table in courses_sheet.tables.values()} == {
        "A2:I4",
        "A7:I8",
    }
    assert courses_sheet.freeze_panes is None
    assert courses_sheet.auto_filter.ref is None

    pastry_sheet = workbook["PASTELERÍA"]
    assert pastry_sheet.max_row == 7
    assert pastry_sheet["A1"].value == (
        "Curso: Pastelería | Horario: Mañana | Grupo: 1 | Inicio: 2026-01-15"
    )
    assert pastry_sheet["A5"].value == (
        "Curso: PASTELERÍA | Horario: Tarde | Grupo: 2 | Inicio: 2026-01-17"
    )
    assert pastry_sheet["D2"].value == "Estudiante"
    assert pastry_sheet["D3"].value == "PÉREZ ANA"
    assert pastry_sheet["E3"].value == "00123456"
    assert pastry_sheet["G3"].value == datetime(2026, 1, 20)
    assert pastry_sheet["K3"].value == "PROMOTOR LOCAL"
    assert pastry_sheet["D6"].value == "Estudiante"
    assert pastry_sheet["D7"].value == "GARCÍA LUIS"
    assert {
        pastry_sheet.cell(row, 4).value for row in (3, 7)
    } == {"PÉREZ ANA", "GARCÍA LUIS"}
    assert {table.ref for table in pastry_sheet.tables.values()} == {
        "A2:M3",
        "A6:M7",
    }
    assert pastry_sheet.auto_filter.ref is None
    assert all(
        pastry_sheet.column_dimensions[column].hidden
        for column in "ABCHIJKLM"
    )
    assert all(
        not pastry_sheet.column_dimensions[column].hidden
        for column in "DEFG"
    )
    bordado_sheet = workbook["BORDADO"]
    assert bordado_sheet["L3"].value == "Sin matrículas"
    workbook.close()


def test_backend_exports_register_to_selected_file(tmp_path):
    application = QCoreApplication.instance()
    if application is None:
        application = QCoreApplication([])
    backend = DesktopBackend(tmp_path / "workspace.sqlite3")
    backend.createWorkspace("Sede inicial", "Av. Principal 123")
    output = tmp_path / "registro.xlsx"

    assert backend.exportRegister(QUrl.fromLocalFile(str(output)).toString())

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["Cursos"]
    workbook.close()
