import re
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

_INVALID_SHEET_CHARACTERS = re.compile(r"[\[\]:*?/\\]")
_INVALID_CELL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_CURRENCY_FORMAT = '"S/ " #,##0.00'
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="173B5E")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SECTION_FILL = PatternFill(fill_type="solid", fgColor="2D587A")
_SECTION_FONT = Font(bold=True, color="FFFFFF")


def export_register(manifest: dict[str, Any], output_path: str | Path) -> Path:
    output = Path(output_path)
    courses = manifest.get("courses", [])
    active_courses = [
        course for course in courses if course.get("active", True)
    ]
    enrollments = manifest.get("enrollments", [])
    students_by_id = {
        student.get("id"): student for student in manifest.get("students", [])
    }
    promoters_by_id = {
        promoter.get("id"): promoter for promoter in manifest.get("promoters", [])
    }
    enrollments_by_course: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    known_course_ids = {course.get("id") for course in courses}
    for enrollment in enrollments:
        course_id = enrollment.get("course_id")
        if course_id in known_course_ids:
            enrollments_by_course[course_id].append(enrollment)

    active_student_ids = {
        student.get("id")
        for student in manifest.get("students", [])
        if student.get("active", True)
    }
    occupied_by_course: dict[Any, set[Any]] = defaultdict(set)
    for enrollment in enrollments:
        student_id = enrollment.get("student_id")
        course_id = enrollment.get("course_id")
        if (
            enrollment.get("active", True)
            and student_id in active_student_ids
            and course_id in known_course_ids
        ):
            occupied_by_course[course_id].add(student_id)

    workbook = Workbook()
    workbook.properties.creator = "EduTrack"
    course_sheet = workbook.active
    course_sheet.title = "Cursos"
    course_headers = [
        "Tipo de curso",
        "Horario",
        "Grupo",
        "Fecha de inicio",
        "Costo de matrícula",
        "Mensualidad",
        "Capacidad",
        "Estudiantes activos",
        "Cupos disponibles",
    ]
    sorted_courses = sorted(
        active_courses,
        key=_course_sort_key,
    )
    grouped_courses: dict[str, list[dict[str, Any]]] = defaultdict(list)
    weekly_groups = _course_week_groups(active_courses)
    week_table_index = 0
    for week_label, week_courses in weekly_groups:
        _write_section_title(course_sheet, week_label, len(course_headers))
        header_row = course_sheet.max_row + 1
        _write_header(course_sheet, course_headers)
        for course in week_courses:
            capacity = _integer(course.get("capacity"))
            active_count = len(occupied_by_course[course.get("id")])
            available = max(0, capacity - active_count)
            course_sheet.append(
                [
                    _text(course.get("name")),
                    _text(course.get("schedule")),
                    _text(course.get("group")),
                    _excel_date(course.get("start_date")),
                    _excel_number(course.get("enrollment_fee")),
                    _excel_number(course.get("monthly_fee")),
                    capacity,
                    active_count,
                    available,
                ]
            )

        week_table_index += 1
        _add_excel_table(
            course_sheet,
            f"RegisterWeekTable{week_table_index}",
            header_row,
            len(course_headers),
        )

    if not weekly_groups:
        _write_header(course_sheet, course_headers)

    for course in sorted_courses:
        key = _text(course.get("name")).strip().casefold() or "sin nombre"
        grouped_courses[key].append(course)

    _format_sheet(
        course_sheet,
        currency_columns=(5, 6),
        date_columns=(4,),
        auto_filter=not weekly_groups,
        freeze_panes="A2" if not weekly_groups else None,
    )

    used_titles = {"cursos"}
    table_index = 0
    for course_group in grouped_courses.values():
        display_name = _text(course_group[0].get("name")).strip() or "Curso"
        title = _unique_sheet_title(display_name, used_titles)
        sheet = workbook.create_sheet(title)
        headers = [
            "Curso",
            "Horario",
            "Grupo",
            "Estudiante",
            "Documento",
            "Teléfonos",
            "Fecha de matrícula",
            "Costo de matrícula",
            "Mensualidad",
            "Código de ficha",
            "Promotor(es)",
            "Estado de matrícula",
            "Estado del curso",
        ]

        for course in course_group:
            _write_section_title(
                sheet,
                _course_table_title(course),
                len(headers),
            )
            header_row = sheet.max_row + 1
            _write_header(sheet, headers)
            course_enrollments = sorted(
                [
                    enrollment
                    for enrollment in enrollments_by_course.get(
                        course.get("id"),
                        [],
                    )
                    if enrollment.get("student_id") in active_student_ids
                ],
                key=lambda enrollment: (
                    _text(enrollment.get("enrollment_date")),
                    _student_name(
                        students_by_id.get(enrollment.get("student_id"))
                    ).casefold(),
                ),
            )
            if not course_enrollments:
                sheet.append(
                    [
                        _text(course.get("name")),
                        _text(course.get("schedule")),
                        _text(course.get("group")),
                        "",
                        "",
                        "",
                        "",
                        None,
                        None,
                        None,
                        "",
                        "Sin matrículas",
                        "Activo" if course.get("active", True) else "Inactivo",
                    ]
                )
            else:
                for enrollment in course_enrollments:
                    student = students_by_id.get(enrollment.get("student_id"))
                    promoter_names = [
                        _text(promoters_by_id[promoter_id].get("name"))
                        for promoter_id in enrollment.get("promoter_ids", [])
                        if promoter_id in promoters_by_id
                    ]
                    if student is None:
                        enrollment_status = "Pendiente de estudiante"
                    else:
                        enrollment_status = (
                            "Activa"
                            if enrollment.get("active", True)
                            else "Inactiva"
                        )
                    sheet.append(
                        [
                            _text(course.get("name")),
                            _text(course.get("schedule")),
                            _text(course.get("group")),
                            _student_name(student) if student else "Pendiente de estudiante",
                            _text(student.get("document_number")) if student else "",
                            ", ".join(student.get("phone_numbers", [])) if student else "",
                            _excel_date(enrollment.get("enrollment_date")),
                            _excel_number(enrollment.get("enrollment_fee")),
                            _excel_number(enrollment.get("monthly_fee")),
                            enrollment.get("record_code"),
                            ", ".join(promoter_names),
                            enrollment_status,
                            "Activo" if course.get("active", True) else "Inactivo",
                        ]
                    )

            table_index += 1
            _add_excel_table(
                sheet,
                f"RegisterCourseTable{table_index}",
                header_row,
                len(headers),
            )

        _format_sheet(
            sheet,
            currency_columns=(8, 9),
            date_columns=(7,),
            auto_filter=False,
            freeze_panes=None,
            hidden_columns=(1, 2, 3, 8, 9, 10, 11, 12, 13),
        )

    workbook.save(output)
    return output


def _course_week_groups(
    courses: list[dict[str, Any]],
) -> list[tuple[str, list[dict[str, Any]]]]:
    courses_by_week: dict[date, list[dict[str, Any]]] = defaultdict(list)
    ungrouped_courses: list[dict[str, Any]] = []
    for course in courses:
        start_date = _text(course.get("start_date")).strip()
        if not _ISO_DATE.fullmatch(start_date):
            ungrouped_courses.append(course)
            continue
        try:
            course_date = date.fromisoformat(start_date)
        except ValueError:
            ungrouped_courses.append(course)
            continue
        week_start = course_date - timedelta(days=course_date.weekday())
        courses_by_week[week_start].append(course)

    groups = []
    for index, week_start in enumerate(sorted(courses_by_week), start=1):
        week_end = week_start + timedelta(days=6)
        label = f"SEMANA {index}: del {_week_range_label(week_start, week_end)}"
        groups.append(
            (
                label,
                sorted(courses_by_week[week_start], key=_course_sort_key),
            )
        )
    if ungrouped_courses:
        groups.append(
            (
                "Semana no definida",
                sorted(ungrouped_courses, key=_course_sort_key),
            )
        )
    return groups


def _course_sort_key(course: dict[str, Any]) -> tuple[str, str, str]:
    return (
        _text(course.get("name")).casefold(),
        _text(course.get("schedule")).casefold(),
        _text(course.get("group")).casefold(),
    )


def _week_range_label(start: date, end: date) -> str:
    months = (
        "enero",
        "febrero",
        "marzo",
        "abril",
        "mayo",
        "junio",
        "julio",
        "agosto",
        "septiembre",
        "octubre",
        "noviembre",
        "diciembre",
    )
    start_month = months[start.month - 1]
    end_month = months[end.month - 1]
    if start.year == end.year and start.month == end.month:
        return f"{start.day} al {end.day} de {end_month} de {end.year}"
    if start.year == end.year:
        return (
            f"{start.day} de {start_month} al {end.day} de "
            f"{end_month} de {end.year}"
        )
    return (
        f"{start.day} de {start_month} de {start.year} al "
        f"{end.day} de {end_month} de {end.year}"
    )


def _add_excel_table(
    sheet: Worksheet,
    display_name: str,
    header_row: int,
    column_count: int,
) -> None:
    table = Table(
        displayName=display_name,
        ref=(
            f"A{header_row}:"
            f"{get_column_letter(column_count)}{sheet.max_row}"
        ),
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    sheet.row_dimensions[sheet.max_row].height = 30


def _write_section_title(
    sheet: Worksheet,
    title: str,
    column_count: int,
) -> None:
    row = (
        1
        if sheet.max_row == 1 and sheet["A1"].value is None
        else sheet.max_row + 2
    )
    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=column_count,
    )
    title_cell = sheet.cell(row=row, column=1, value=title)
    title_cell.fill = _SECTION_FILL
    title_cell.font = _SECTION_FONT
    title_cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in range(2, column_count + 1):
        sheet.cell(row=row, column=column).fill = _SECTION_FILL
    sheet.row_dimensions[row].height = 30


def _course_table_title(course: dict[str, Any]) -> str:
    name = _text(course.get("name")).strip() or "Sin nombre"
    schedule = _text(course.get("schedule")).strip() or "Sin horario"
    group = _text(course.get("group")).strip() or "Sin grupo"
    parts = [
        f"Curso: {name}",
        f"Horario: {schedule}",
        f"Grupo: {group}",
    ]
    start_date = _text(course.get("start_date")).strip()
    if start_date:
        parts.append(f"Inicio: {start_date}")
    return " | ".join(parts)


def _format_sheet(
    sheet: Worksheet,
    currency_columns: tuple[int, ...],
    date_columns: tuple[int, ...],
    *,
    auto_filter: bool = True,
    freeze_panes: str | None = "A2",
    hidden_columns: tuple[int, ...] = (),
) -> None:
    if freeze_panes is not None:
        sheet.freeze_panes = freeze_panes
    if auto_filter:
        sheet.auto_filter.ref = sheet.dimensions
    for column in hidden_columns:
        sheet.column_dimensions[get_column_letter(column)].hidden = True

    for column in currency_columns:
        for cell in sheet.iter_cols(
            min_col=column,
            max_col=column,
            min_row=2,
        ):
            for value in cell:
                value.number_format = _CURRENCY_FORMAT
    for column in date_columns:
        for cell in sheet.iter_cols(
            min_col=column,
            max_col=column,
            min_row=2,
        ):
            for value in cell:
                if isinstance(value.value, date):
                    value.number_format = "dd/mm/yyyy"

    for column_index, cells in enumerate(sheet.columns, start=1):
        max_length = max(
            (len(str(cell.value)) for cell in cells if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 12),
            40,
        )


def _unique_sheet_title(name: str, used_titles: set[str]) -> str:
    base = _INVALID_SHEET_CHARACTERS.sub("-", name.strip().strip("'")).upper()
    base = base[:31] or "CURSO"
    candidate = base
    suffix = 2
    while candidate.casefold() in used_titles:
        suffix_text = f" ({suffix})"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_titles.add(candidate.casefold())
    return candidate


def _text(value: Any) -> str:
    if value is None:
        return ""
    return _INVALID_CELL_CHARACTERS.sub("", str(value))


def _student_name(student: dict[str, Any] | None) -> str:
    if not student:
        return ""
    return " ".join(
        part
        for part in (
            _text(student.get("last_name")).strip(),
            _text(student.get("first_name")).strip(),
        )
        if part
    )


def _excel_date(value: Any) -> date | str:
    text = _text(value)
    if not text:
        return ""
    try:
        return date.fromisoformat(text)
    except ValueError:
        return text


def _excel_number(value: Any) -> Decimal | str:
    text = _text(value).strip()
    if not text:
        return ""
    try:
        return Decimal(text)
    except InvalidOperation:
        return text


def _integer(value: Any) -> int:
    try:
        return max(0, int(value))
    except (TypeError, ValueError):
        return 0
