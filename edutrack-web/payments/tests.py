import threading
from datetime import date, datetime
from io import BytesIO
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.db import IntegrityError, connections, transaction
from django.db.models import ProtectedError, Sum
from django.test import Client, RequestFactory, TestCase, TransactionTestCase, skipUnlessDBFeature
from django.urls import reverse
from openpyxl import load_workbook
from django.utils import timezone

from core.models import Campus
from core.testing import create_user, grant, revoke
from payments.forms import PaymentForm, PaymentUpdateForm
from payments.models import Installment, Payment, PaymentAllocation
from students.models import Course, Enrollment, Student


class PaymentCrudTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede de pagos", active=True)
        self.user = create_user(
            username="cashier",
            password="secret123",
            manage_students=True,
            register_payments=True,
        )
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        student = Student.objects.create(campus=self.campus, full_name="Alumno para pago")
        course = Course.objects.create(
            campus=self.campus,
            name="Curso para pago",
            start_date=date(2026, 10, 14),
            enrollment_fee=40,
            monthly_fee=60,
        )
        self.enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=student,
            course=course,
            enrollment_fee=40,
            monthly_fee=60,
        )
        self.installment = Installment.objects.create(
            enrollment=self.enrollment,
            concept="Cuota inicial",
            amount=120,
            due_date=date(2026, 10, 14),
        )
        self.payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="D1",
            amount=40,
        )
        self.allocation = PaymentAllocation.objects.create(
            payment=self.payment,
            installment=self.installment,
            amount=40,
        )

    def test_dashboard_totals_match_installment_balances(self):
        Installment.objects.create(
            enrollment=self.enrollment, concept="Vencida", amount=50, due_date=date(2020, 1, 1)
        )
        Installment.objects.create(
            enrollment=self.enrollment, concept="Futura", amount=30, due_date=date(2999, 1, 1)
        )
        self.user.is_superuser = True
        self.user.save()
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        today = timezone.localdate()
        installments = list(Installment.objects.filter(enrollment=self.enrollment))
        self.assertEqual(response.context["pending_total"], sum(i.balance for i in installments))
        self.assertEqual(
            response.context["overdue_count"],
            sum(1 for i in installments if i.balance > 0 and i.due_date < today),
        )
        self.assertEqual(response.context["overdue_count"], 2)

    def test_dashboard_stats_refresh_after_a_new_payment(self):
        self.user.is_superuser = True
        self.user.save()
        first = self.client.get(reverse("dashboard")).context["payments_total"]
        payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="D77",
            amount=25,
        )
        PaymentAllocation.objects.create(payment=payment, installment=self.installment, amount=25)
        second = self.client.get(reverse("dashboard")).context
        self.assertEqual(second["payments_total"], first + 25)
        self.assertEqual(second["pending_total"], Decimal("55.00"))

    def test_payment_list_exports_filtered_excel(self):
        response = self.client.get(reverse("payment-excel"), {"receipt": "D1"})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        values = [cell.value for row in sheet.iter_rows(min_row=5) for cell in row]
        self.assertIn("D1", values)
        none = self.client.get(reverse("payment-excel"), {"receipt": "NOEXISTE"})
        sheet = load_workbook(BytesIO(none.content)).active
        self.assertNotIn("D1", [c.value for row in sheet.iter_rows(min_row=5) for c in row])

    def test_payment_export_requires_export_permission(self):
        revoke(self.user, "manage_students")
        self.assertEqual(self.client.get(reverse("payment-excel")).status_code, 403)

    def test_cash_closing_groups_by_cashier_and_ignores_voided_in_total(self):
        other = create_user(username="cajera2", password="x" * 12, register_payments=True)
        other.campuses.add(self.campus)
        Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=other,
            receipt_number="D2",
            amount=15,
            status=Payment.Status.VOID,
        )
        response = self.client.get(reverse("cash-closing"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_confirmed"], Decimal("40"))
        self.assertEqual(response.context["total_voided"], Decimal("15"))
        self.assertEqual(len(response.context["groups"]), 2)
        only = self.client.get(reverse("cash-closing"), {"cajero": other.pk})
        self.assertEqual(only.context["total_confirmed"], Decimal("0"))
        past = self.client.get(reverse("cash-closing"), {"fecha": "2000-01-01"})
        self.assertEqual(past.context["groups"], [])

    def test_signed_cash_closing_freezes_totals_and_locks_payments(self):
        from payments.models import CashClosing

        today = timezone.localdate().strftime("%Y-%m-%d")
        self.assertEqual(
            self.client.post(reverse("cash-closing-close"), {"fecha": today}).status_code, 403
        )
        grant(self.user, "core.close_cash")
        self.client.post(reverse("cash-closing-close"), {"fecha": today, "nota": "ok"})
        closing = CashClosing.objects.get(campus=self.campus)
        self.assertEqual(closing.confirmed_total, Decimal("40"))
        self.assertEqual(closing.closed_by, self.user)

        edit = self.client.post(reverse("payment-update", args=[self.payment.pk]), {})
        self.assertRedirects(edit, reverse("payment-list"), fetch_redirect_response=False)
        self.client.post(reverse("payment-delete", args=[self.payment.pk]))
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.status, Payment.Status.CONFIRMED)
        # Cerrar de nuevo no duplica el cierre
        self.client.post(reverse("cash-closing-close"), {"fecha": today})
        self.assertEqual(CashClosing.objects.count(), 1)

        # Solo un superusuario reabre
        reopen = self.client.post(reverse("cash-closing-reopen"), {"cierre": closing.pk})
        self.assertEqual(reopen.status_code, 403)
        self.user.is_superuser = True
        self.user.save()
        self.client.post(reverse("cash-closing-reopen"), {"cierre": closing.pk})
        self.assertFalse(CashClosing.objects.exists())

    def test_cash_closing_excel_and_permissions(self):
        response = self.client.get(reverse("cash-closing-excel"))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertIn("D1", [c.value for row in sheet.iter_rows(min_row=4) for c in row])
        revoke(self.user, "register_payments")
        self.assertEqual(self.client.get(reverse("cash-closing")).status_code, 403)

    def test_pending_debts_export_and_overdue_filter(self):
        Installment.objects.create(
            enrollment=self.enrollment, concept="Futura", amount=30, due_date=date(2999, 1, 1)
        )
        response = self.client.get(reverse("pending-debts-excel"))
        sheet = load_workbook(BytesIO(response.content)).active
        concepts = [row[3].value for row in sheet.iter_rows(min_row=5)]
        self.assertIn("Futura", concepts)
        overdue = self.client.get(reverse("pending-debts"), {"vencidas": "1"})
        names = [d.concept for d in overdue.context["debts"]]
        self.assertIn("Cuota inicial", names)
        self.assertNotIn("Futura", names)

    def test_enrollment_statement_shows_installments_and_payments(self):
        response = self.client.get(reverse("enrollment-statement", args=[self.enrollment.pk]))
        self.assertContains(response, "ESTADO DE CUENTA")
        self.assertContains(response, "Cuota inicial")
        self.assertContains(response, "D1")
        self.assertEqual(response.context["total_balance"], Decimal("80.00"))

    def test_enrollment_statement_is_limited_to_the_selected_campus(self):
        other = Campus.objects.create(name="Otra sede", active=True)
        student = Student.objects.create(campus=other, full_name="Alumno ajeno")
        foreign = Enrollment.objects.create(campus=other, student=student)
        response = self.client.get(reverse("enrollment-statement", args=[foreign.pk]))
        self.assertEqual(response.status_code, 404)

    def test_payment_views_are_limited_to_the_selected_campus(self):
        other = Campus.objects.create(name="Otra sede para pagos", active=True)
        student = Student.objects.create(campus=other, full_name="Alumno ajeno")
        enrollment = Enrollment.objects.create(campus=other, student=student)
        payment = Payment.objects.create(
            campus=other,
            enrollment=enrollment,
            cashier=self.user,
            receipt_number="AJENO-1",
            amount=10,
        )

        for url, method in (
            (reverse("payment-update", args=[payment.pk]), "get"),
            (reverse("payment-delete", args=[payment.pk]), "get"),
            (reverse("payment-receipt", args=[payment.pk]), "get"),
            (reverse("payment-delete", args=[payment.pk]), "post"),
        ):
            with self.subTest(url=url, method=method):
                self.assertEqual(getattr(self.client, method)(url).status_code, 404)

    def test_payment_views_require_their_permissions(self):
        revoke(self.user, "register_payments")

        self.assertEqual(self.client.get(reverse("payment-list")).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("payment-create-for", args=[self.enrollment.pk])).status_code,
            403,
        )

    def test_partial_payment_updates_balance_and_void_restores_it(self):
        url = reverse("payment-create-for", args=[self.enrollment.pk])
        response = self.client.post(
            url,
            {
                "receipt_number_0": "D",
                "receipt_number_1": "2",
                "enrollment": self.enrollment.pk,
                "amount": "50.00",
                "reference": "",
                "notes": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        payment = Payment.objects.get(receipt_number="D2")
        self.installment.refresh_from_db()
        self.assertEqual(self.installment.balance, Decimal("30.00"))
        allocation = PaymentAllocation.objects.get(payment=payment)
        self.assertEqual(allocation.amount, Decimal("50.00"))

        response = self.client.post(reverse("payment-delete", args=[payment.pk]))
        self.assertEqual(response.status_code, 302)
        payment.refresh_from_db()
        self.installment.refresh_from_db()
        self.assertEqual(payment.status, Payment.Status.VOID)
        self.assertEqual(self.installment.balance, Decimal("80.00"))
        self.assertTrue(PaymentAllocation.objects.filter(pk=allocation.pk).exists())

    def test_voucher_is_served_only_to_users_of_its_campus(self):
        self.payment.voucher = "vouchers/prueba.png"
        self.payment.save()
        url = "/media/vouchers/prueba.png"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Accel-Redirect"], "/_protected_media/vouchers/prueba.png")
        outsider = create_user(username="outsider", password="x" * 12)
        self.client.force_login(outsider)
        self.assertEqual(self.client.get(url).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_media_path_rejects_traversal(self):
        self.assertEqual(self.client.get("/media/vouchers/../../etc/passwd").status_code, 404)

    def test_payment_can_be_updated_and_reallocated(self):
        response = self.client.get(reverse("payment-update", args=[self.payment.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(
            reverse("payment-update", args=[self.payment.pk]),
            {
                "receipt_number_0": "D",
                "receipt_number_1": "2",
                "enrollment": self.enrollment.pk,
                "amount": "70.00",
                "reference": "Transferencia",
                "notes": "Actualizado",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.receipt_number, "D2")
        self.assertEqual(self.payment.amount, Decimal("70.00"))
        self.assertEqual(
            PaymentAllocation.objects.filter(payment=self.payment).aggregate(total=Sum("amount"))[
                "total"
            ],
            Decimal("70.00"),
        )

    def test_payment_receipt_is_saved_with_uppercase_prefix(self):
        payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="d100",
            amount=10,
        )

        self.assertEqual(payment.receipt_number, "D100")
        payment.refresh_from_db()
        self.assertEqual(payment.receipt_number, "D100")

    def test_receipt_form_uses_optional_letter_and_positive_integer_inputs(self):
        response = self.client.get(reverse("payment-create-for", args=[self.enrollment.pk]))
        self.assertContains(response, 'name="receipt_number_0"')
        self.assertContains(response, 'name="receipt_number_1"')
        self.assertContains(response, "data-uppercase")

        request = RequestFactory().get(reverse("payment-create"))
        request.session = {"selected_campus_id": self.campus.pk}
        request.user = self.user
        cases = [
            ({"receipt_number_0": "", "receipt_number_1": "55"}, "55"),
            ({"receipt_number_0": "d", "receipt_number_1": "100"}, "D100"),
            ({"receipt_number_0": "AB", "receipt_number_1": "5"}, None),
            ({"receipt_number_0": "1", "receipt_number_1": "5"}, None),
            ({"receipt_number_0": "D", "receipt_number_1": "12a"}, None),
            ({"receipt_number_0": "D", "receipt_number_1": "-5"}, None),
            ({"receipt_number_0": "D", "receipt_number_1": "0"}, None),
            ({"receipt_number_0": "D", "receipt_number_1": ""}, None),
        ]
        for data, expected in cases:
            with self.subTest(data=data):
                form = PaymentForm(data=data, user=self.user, request=request)
                form.is_valid()
                if expected is None:
                    self.assertIn("receipt_number", form.errors)
                else:
                    self.assertNotIn("receipt_number", form.errors)
                    self.assertEqual(form.cleaned_data["receipt_number"], expected)

    def test_receipt_number_must_be_unique_within_campus_only(self):
        request = RequestFactory().get(reverse("payment-create"))
        request.session = {"selected_campus_id": self.campus.pk}
        request.user = self.user
        same_campus_form = PaymentForm(
            data={"receipt_number_0": "D", "receipt_number_1": "1"},
            user=self.user,
            request=request,
        )
        self.assertFalse(same_campus_form.is_valid())
        self.assertIn("receipt_number", same_campus_form.errors)

        other_campus = Campus.objects.create(name="Otra sede", active=True)
        self.user.campuses.add(other_campus)
        request = RequestFactory().get(reverse("payment-create"))
        request.session = {"selected_campus_id": other_campus.pk}
        request.user = self.user
        other_campus_form = PaymentForm(
            data={"receipt_number_0": "D", "receipt_number_1": "1"},
            user=self.user,
            request=request,
        )
        self.assertFalse(other_campus_form.is_valid())
        self.assertNotIn("receipt_number", other_campus_form.errors)

        other_payment = Payment.objects.create(
            campus=other_campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="D1",
            amount=10,
        )
        self.assertEqual(other_payment.receipt_number, self.payment.receipt_number)

        request = RequestFactory().get(reverse("payment-update", args=[self.payment.pk]))
        request.session = {"selected_campus_id": self.campus.pk}
        request.user = self.user
        update_form = PaymentUpdateForm(
            data={
                "receipt_number_0": "D",
                "receipt_number_1": "1",
                "enrollment": self.enrollment.pk,
                "amount": "40.00",
            },
            instance=self.payment,
            user=self.user,
            request=request,
        )
        self.assertTrue(update_form.is_valid(), update_form.errors)

        with self.assertRaises(IntegrityError), transaction.atomic():
            Payment.objects.create(
                campus=self.campus,
                enrollment=self.enrollment,
                cashier=self.user,
                receipt_number="D1",
                amount=10,
            )

    def test_payment_list_filters_and_sorts_receipts_students_and_dates(self):
        other_student = Student.objects.create(
            campus=self.campus,
            full_name="Beatriz de pagos",
        )
        other_course = Course.objects.create(
            campus=self.campus,
            name="Otro curso",
            start_date=date(2026, 10, 14),
        )
        other_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=other_student,
            course=other_course,
        )
        other_payment = Payment.objects.create(
            campus=self.campus,
            enrollment=other_enrollment,
            cashier=self.user,
            receipt_number="D2",
            amount=20,
        )
        void_payment = Payment.objects.create(
            campus=self.campus,
            enrollment=other_enrollment,
            cashier=self.user,
            receipt_number="D3",
            amount=10,
            status=Payment.Status.VOID,
        )
        Payment.objects.filter(pk=self.payment.pk).update(
            created_at=timezone.make_aware(datetime(2026, 10, 10, 9, 0))
        )
        Payment.objects.filter(pk=other_payment.pk).update(
            created_at=timezone.make_aware(datetime(2026, 10, 11, 9, 0))
        )
        Payment.objects.filter(pk=void_payment.pk).update(
            created_at=timezone.make_aware(datetime(2026, 10, 12, 9, 0))
        )

        response = self.client.get(reverse("payment-list"), {"sort": "receipt", "direction": "asc"})
        self.assertEqual(response.context["payments"][0].receipt_number, "D1")
        self.assertContains(response, 'name="receipt"')
        self.assertContains(response, 'name="date_from"')
        self.assertContains(response, 'name="date_to"')
        self.assertContains(response, 'for="payment-date-from">Desde')
        self.assertContains(response, 'for="payment-date-to">Hasta')
        self.assertNotContains(response, "<th>Sede</th>")
        self.assertContains(response, "?sort=receipt&amp;direction=desc")
        self.assertEqual(response.context["payments"][0].row_number, 1)

        student_sort = self.client.get(
            reverse("payment-list"), {"sort": "student", "direction": "asc"}
        )
        self.assertEqual(student_sort.context["payments"][0].pk, self.payment.pk)
        date_sort = self.client.get(reverse("payment-list"), {"sort": "date", "direction": "asc"})
        self.assertEqual(date_sort.context["payments"][0].pk, self.payment.pk)

        for query, expected_payments in (
            ({"student": "Beatriz"}, [other_payment, void_payment]),
            ({"receipt": "D1"}, self.payment),
            ({"status": "ANULADO"}, void_payment),
            ({"date_from": "2026-10-10", "date_to": "2026-10-10"}, self.payment),
            ({"date_from": "2026-10-13"}, None),
        ):
            with self.subTest(query=query):
                response = self.client.get(reverse("payment-list"), query)
                if expected_payments is None:
                    self.assertEqual(list(response.context["payments"]), [])
                else:
                    if not isinstance(expected_payments, list):
                        expected_payments = [expected_payments]
                    self.assertEqual(
                        {payment.pk for payment in response.context["payments"]},
                        {payment.pk for payment in expected_payments},
                    )

    def test_payment_list_sorts_receipts_by_type_letter_and_number(self):
        for receipt_number in ["44", "1", "7", "D45", "D23", "A43", "A34"]:
            Payment.objects.create(
                campus=self.campus,
                enrollment=self.enrollment,
                cashier=self.user,
                receipt_number=receipt_number,
                amount=1,
            )

        response = self.client.get(reverse("payment-list"), {"sort": "receipt", "direction": "asc"})

        self.assertEqual(
            [payment.receipt_number for payment in response.context["payments"]],
            ["1", "7", "44", "A34", "A43", "D1", "D23", "D45"],
        )

    def test_receipt_check_lists_missing_numeric_receipts(self):
        for receipt_number in ["1", "2", "4", "5"]:
            Payment.objects.create(
                campus=self.campus,
                enrollment=self.enrollment,
                cashier=self.user,
                receipt_number=receipt_number,
                amount=1,
            )

        response = self.client.get(
            reverse("tools"),
            {"receipt_start": "1", "receipt_end": "7", "check_receipts": "1"},
        )

        self.assertEqual(response.context["missing_receipts"], ["3", "6", "7"])
        for number in (3, 6, 7):
            self.assertContains(response, f'<span class="tool-chip">{number}</span>')

    def test_receipt_check_lists_missing_receipts_with_the_same_letter(self):
        for receipt_number in ["D2", "D4", "D5"]:
            Payment.objects.create(
                campus=self.campus,
                enrollment=self.enrollment,
                cashier=self.user,
                receipt_number=receipt_number,
                amount=1,
            )

        response = self.client.get(
            reverse("tools"),
            {"receipt_start": "D1", "receipt_end": "D4", "check_receipts": "1"},
        )

        self.assertEqual(response.context["missing_receipts"], ["D3"])

    def test_receipt_check_rejects_different_letter_prefixes(self):
        response = self.client.get(
            reverse("tools"),
            {"receipt_start": "D1", "receipt_end": "E4", "check_receipts": "1"},
        )

        self.assertIsNone(response.context["missing_receipts"])
        self.assertContains(response, "usar la misma letra")

    def test_voucher_is_optional_and_saved_as_jpeg(self):
        import io
        import shutil
        import tempfile

        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from PIL import Image

        buffer = io.BytesIO()
        Image.new("RGBA", (30, 30), (255, 0, 0, 128)).save(buffer, "PNG")
        url = reverse("payment-update", args=[self.payment.pk])
        data = {
            "receipt_number_0": "D",
            "receipt_number_1": "2",
            "enrollment": self.enrollment.pk,
            "amount": "70.00",
        }
        media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media, ignore_errors=True)
        with override_settings(MEDIA_ROOT=media):
            self.assertEqual(self.client.post(url, data).status_code, 302)
            self.payment.refresh_from_db()
            self.assertFalse(self.payment.voucher)

            upload = SimpleUploadedFile("v.png", buffer.getvalue(), content_type="image/png")
            with self.captureOnCommitCallbacks(execute=True):
                response = self.client.post(url, {**data, "voucher": upload})
            self.assertEqual(response.status_code, 302)
            self.payment.refresh_from_db()
            self.assertTrue(self.payment.voucher.name.startswith("vouchers/"))
            self.assertTrue(self.payment.voucher.name.endswith(".jpg"))
            with self.payment.voucher.open("rb") as handle:
                self.assertEqual(Image.open(handle).format, "JPEG")

            self.client.post(url, data)
            self.payment.refresh_from_db()
            self.assertTrue(self.payment.voucher)

            with self.captureOnCommitCallbacks(execute=True):
                self.client.post(url, {**data, "voucher-clear": "on"})
            self.payment.refresh_from_db()
            self.assertFalse(self.payment.voucher)

    def test_payment_can_be_voided_without_deleting_receipt_or_allocation(self):
        response = self.client.get(reverse("payment-delete", args=[self.payment.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(reverse("payment-delete", args=[self.payment.pk]), follow=True)
        self.assertEqual(response.status_code, 200)
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.status, Payment.Status.VOID)
        self.assertTrue(Payment.objects.filter(pk=self.payment.pk).exists())
        self.assertTrue(PaymentAllocation.objects.filter(pk=self.allocation.pk).exists())

    def test_select_view_extra_filters_narrow_pending_enrollments(self):
        url = reverse("payment-create")

        def rows(params):
            return list(self.client.get(url, params).context["enrollments"])

        self.assertEqual(len(rows({})), 1)
        self.assertEqual(len(rows({"course": self.enrollment.course_id})), 1)
        self.assertEqual(len(rows({"course": 999999})), 0)
        self.assertEqual(len(rows({"debt_min": "0.01", "debt_max": "100000"})), 1)
        self.assertEqual(len(rows({"debt_min": "100000"})), 0)
        self.assertEqual(len(rows({"date_from": "2999-01-01"})), 0)
        self.assertEqual(len(rows({"ficha": "abc"})), 0)

    def test_global_search_does_not_return_records(self):
        url = reverse("global-search")
        for term in ("alumno para", "d1"):
            response = self.client.get(url, {"q": term})
            self.assertNotContains(response, "ALUMNO PARA PAGO")
            self.assertNotContains(response, "Ver recibo")
        self.assertContains(self.client.get(url, {"q": "zzzz"}), "No se encontraron secciones")

    def test_global_search_hides_sections_without_permission(self):
        revoke(self.user, "register_payments")
        response = self.client.get(reverse("global-search"), {"q": "pagos"})
        self.assertNotContains(response, f'href="{reverse("payment-list")}"')

    def test_cannot_register_payment_when_todays_cash_is_closed(self):
        from payments.models import CashClosing

        CashClosing.objects.create(
            campus=self.campus,
            date=timezone.localdate(),
            confirmed_total=0,
            voided_total=0,
            payments_count=0,
            closed_by=self.user,
        )
        before = Payment.objects.count()
        response = self.client.post(
            reverse("payment-create-for", args=[self.enrollment.pk]), {"amount": "40"}
        )
        self.assertRedirects(response, reverse("payment-list"), fetch_redirect_response=False)
        self.assertEqual(Payment.objects.count(), before)

    def test_cash_closing_rejects_future_dates_and_hides_button_without_permission(self):
        self.assertNotContains(self.client.get(reverse("cash-closing")), "Cerrar caja del")
        grant(self.user, "core.close_cash")
        self.assertContains(self.client.get(reverse("cash-closing")), "Cerrar caja del")
        response = self.client.post(reverse("cash-closing-close"), {"fecha": "2999-01-01"})
        self.assertRedirects(response, reverse("cash-closing"))
        response = self.client.post(reverse("cash-closing-close"), {"fecha": "no-es-fecha"})
        self.assertRedirects(response, reverse("cash-closing"))
        self.assertEqual(self.client.get(reverse("cash-closing-close")).status_code, 405)

    def test_global_search_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("global-search"), {"q": "x"})
        self.assertEqual(response.status_code, 302)

    def test_select_view_lists_pending_enrollments_with_pay_column(self):
        response = self.client.get(reverse("payment-create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Pagar")
        self.assertContains(response, reverse("payment-create-for", args=[self.enrollment.pk]))
        self.assertContains(response, "S/ 80,00")

    def test_legacy_enrollment_param_redirects_to_payment_form(self):
        response = self.client.get(reverse("payment-create"), {"enrollment": self.enrollment.pk})
        self.assertRedirects(response, reverse("payment-create-for", args=[self.enrollment.pk]))

    def test_payment_form_has_enrollment_loaded(self):
        response = self.client.get(reverse("payment-create-for", args=[self.enrollment.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Deuda pendiente")
        form = response.context["form"]
        self.assertEqual(str(form["enrollment"].value()), str(self.enrollment.pk))
        self.assertEqual(Decimal(str(form["amount"].value())), Decimal("80"))

    def test_payment_form_rejects_enrollment_without_debt(self):
        response = self.client.get(reverse("payment-create-for", args=[999999]))
        self.assertRedirects(response, reverse("payment-create"))


class PaymentListExtraFiltersTests(TestCase):
    def test_amount_and_cashier_filters(self):
        from decimal import Decimal
        from core.models import Campus
        from payments.models import Payment
        from students.models import Course, Enrollment, Student

        campus = Campus.objects.create(name="Sede Filtros", active=True)
        user = create_user(username="caja1", password="pass-12345", register_payments=True)
        user.campuses.add(campus)
        student = Student.objects.create(campus=campus, full_name="Ana", phone="987654321")
        course = Course.objects.create(campus=campus, name="Ingles", schedule="L", group="A")
        enrollment = Enrollment.objects.create(campus=campus, student=student, course=course)
        for number, amount in (("1", "10"), ("2", "50")):
            Payment.objects.create(
                campus=campus,
                enrollment=enrollment,
                cashier=user,
                receipt_number=number,
                amount=Decimal(amount),
            )
        self.client.force_login(user)
        session = self.client.session
        session["selected_campus_id"] = campus.pk
        session.save()
        url = reverse("payment-list")
        response = self.client.get(url, {"amount_min": "20", "cashier": "caja"})
        self.assertEqual([p.receipt_number for p in response.context["payments"]], ["2"])
        response = self.client.get(url, {"amount_max": "abc"})
        self.assertEqual(len(response.context["payments"]), 0)


class PaymentAllocationIntegrityTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede integridad", active=True)
        self.user = create_user(username="cashier2", password="x" * 12, register_payments=True)
        course = Course.objects.create(
            campus=self.campus, name="Curso", start_date=date(2026, 10, 14)
        )
        self.enrollments = [
            Enrollment.objects.create(
                campus=self.campus,
                student=Student.objects.create(campus=self.campus, full_name=f"Alumno {n}"),
                course=course,
            )
            for n in range(2)
        ]
        self.installment = Installment.objects.create(
            enrollment=self.enrollments[0], concept="Cuota", amount=100, due_date=date(2026, 10, 14)
        )
        self.payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollments[0],
            cashier=self.user,
            receipt_number="X1",
            amount=60,
        )

    def test_rejects_installment_from_another_enrollment(self):
        other = Installment.objects.create(
            enrollment=self.enrollments[1], concept="Otra", amount=100, due_date=date(2026, 10, 14)
        )
        with self.assertRaises(ValidationError):
            PaymentAllocation.objects.create(payment=self.payment, installment=other, amount=10)

    def test_rejects_allocation_above_payment_amount(self):
        PaymentAllocation.objects.create(
            payment=self.payment, installment=self.installment, amount=50
        )
        with self.assertRaises(ValidationError):
            PaymentAllocation.objects.create(
                payment=self.payment, installment=self.installment, amount=11
            )

    def test_rejects_allocation_above_installment_amount(self):
        PaymentAllocation.objects.create(
            payment=self.payment, installment=self.installment, amount=60
        )
        second = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollments[0],
            cashier=self.user,
            receipt_number="X2",
            amount=60,
        )
        with self.assertRaises(ValidationError):
            PaymentAllocation.objects.create(
                payment=second, installment=self.installment, amount=41
            )

    def test_database_rejects_non_positive_allocation(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            PaymentAllocation.objects.bulk_create(
                [PaymentAllocation(payment=self.payment, installment=self.installment, amount=0)]
            )

    def test_payments_cannot_be_deleted(self):
        with self.assertRaises(ProtectedError):
            self.payment.delete()
        self.assertTrue(Payment.objects.filter(pk=self.payment.pk).exists())


@skipUnlessDBFeature("has_select_for_update")
class ConcurrentPaymentRegistrationTests(TransactionTestCase):
    serialized_rollback = True

    def test_two_cashiers_cannot_collect_more_than_the_remaining_balance(self):
        campus = Campus.objects.create(name="Sede pagos concurrentes", active=True)
        course = Course.objects.create(
            campus=campus,
            name="Curso concurrente",
            start_date=date(2026, 10, 14),
            enrollment_fee=40,
            monthly_fee=60,
        )
        student = Student.objects.create(campus=campus, full_name="Alumno concurrente")
        enrollment = Enrollment.objects.create(
            campus=campus,
            student=student,
            course=course,
            enrollment_fee=40,
            monthly_fee=60,
        )
        installment = Installment.objects.create(
            enrollment=enrollment,
            concept="Cuota concurrente",
            amount=100,
            due_date=date(2026, 10, 14),
        )
        cashiers = [
            create_user(username=f"cajero-{index}", register_payments=True) for index in range(2)
        ]
        for cashier in cashiers:
            cashier.campuses.add(campus)

        barrier = threading.Barrier(2)
        results = []
        errors = []

        def submit(cashier, receipt_number):
            connections.close_all()
            try:
                client = Client()
                client.force_login(cashier)
                session = client.session
                session["selected_campus_id"] = campus.pk
                session.save()
                barrier.wait(timeout=10)
                response = client.post(
                    reverse("payment-create-for", args=[enrollment.pk]),
                    {
                        "receipt_number_0": "C",
                        "receipt_number_1": receipt_number,
                        "enrollment": enrollment.pk,
                        "amount": "60.00",
                        "reference": "",
                        "notes": "",
                    },
                )
                results.append(response.status_code)
            except Exception as error:
                errors.append(error)
            finally:
                connections.close_all()

        threads = [
            threading.Thread(target=submit, args=(cashiers[0], "1")),
            threading.Thread(target=submit, args=(cashiers[1], "2")),
        ]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)

        self.assertFalse(any(thread.is_alive() for thread in threads), "Payment request hung.")
        self.assertEqual(errors, [])
        self.assertCountEqual(results, [200, 302])
        self.assertEqual(
            Payment.objects.filter(enrollment=enrollment, status=Payment.Status.CONFIRMED).count(),
            1,
        )
        self.assertEqual(
            PaymentAllocation.objects.filter(payment__enrollment=enrollment).aggregate(
                total=Sum("amount")
            )["total"],
            Decimal("60.00"),
        )
        installment.refresh_from_db()
        self.assertEqual(installment.balance, Decimal("40.00"))
