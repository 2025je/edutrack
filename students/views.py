from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import (
    Case,
    Count,
    DecimalField,
    F,
    Max,
    Min,
    OuterRef,
    Prefetch,
    Q,
    Subquery,
    Sum,
    Value,
    When,
)
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.http import JsonResponse
from django.urls import reverse_lazy
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.generic import ListView, CreateView, UpdateView, DeleteView, TemplateView, View
from itertools import groupby
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from io import BytesIO
from core.campus import SelectedCampusRequiredMixin, FeaturePermissionMixin, get_selected_campus
from .models import (
    EMPTY_VALUE,
    Student,
    Enrollment,
    Course,
    Ficha,
)
from .forms import StudentForm, StudentPhoneFormSet, EnrollmentForm
from .installments import fee_conflicts, sync_initial_installments
from .search import filter_students


MAX_RECORD_CODE_CHECK_RANGE = 10_000


class CampusScopedMixin(FeaturePermissionMixin, SelectedCampusRequiredMixin, LoginRequiredMixin):
    required_permission = "students.view_student"

    def campus_queryset(self, qs, field="campus"):
        campus = get_selected_campus(self.request)
        return qs.filter(**{field: campus})


class StudentListView(CampusScopedMixin, ListView):
    model = Student
    template_name = "students/student_list.html"
    context_object_name = "students"
    paginate_by = 25
    sortable_fields = {
        "row_number": "created_at",
        "name": "full_name",
        "enrollments": "enrollment_count",
    }

    def get_queryset(self):
        qs = self.campus_queryset(
            Student.objects.select_related("campus")
            .annotate(
                enrollment_count=Count(
                    "enrollments",
                    filter=Q(enrollments__active=True, enrollments__course__isnull=False),
                    distinct=True,
                ),
                unassigned_count=Count(
                    "enrollments",
                    filter=Q(enrollments__active=True, enrollments__course__isnull=True),
                    distinct=True,
                ),
            )
            .order_by("full_name", "pk")
        )
        q = self.request.GET.get("q", "").strip()
        if q:
            qs = filter_students(qs, q, with_fichas=True)
        sort = self.request.GET.get("sort", "row_number")
        direction = self.request.GET.get("direction", "desc")
        sort_field = self.sortable_fields.get(sort, "full_name")
        if direction not in {"asc", "desc"}:
            direction = "desc"
        if sort == "row_number":
            chronological_fields = ["created_at", "pk"]
            if direction == "desc":
                chronological_fields = [f"-{field}" for field in chronological_fields]
            return qs.order_by(*chronological_fields)
        if direction == "desc":
            sort_field = f"-{sort_field}"
        return qs.order_by(sort_field, "full_name", "pk")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["payable_student_ids"] = set()
        if self.request.user.has_perm("payments.add_payment"):
            from payments.forms import pending_enrollments_for_campus

            context["payable_student_ids"] = set(
                pending_enrollments_for_campus(get_selected_campus(self.request)).values_list(
                    "student_id", flat=True
                )
            )
        context["current_sort"] = self.request.GET.get("sort", "row_number")
        context["current_direction"] = self.request.GET.get("direction", "desc")
        context["sort_urls"] = {}
        for key in self.sortable_fields:
            sort_params = self.request.GET.copy()
            sort_params.pop("page", None)
            if context["current_sort"] == key:
                next_direction = "desc" if context["current_direction"] == "asc" else "asc"
            else:
                next_direction = "desc" if key == "row_number" else "asc"
            sort_params["sort"] = key
            sort_params["direction"] = next_direction
            context["sort_urls"][key] = f"?{sort_params.urlencode()}"
        chronological_ids = self.object_list.order_by("created_at", "pk").values_list(
            "pk", flat=True
        )
        row_numbers = {
            student_id: row_number
            for row_number, student_id in enumerate(chronological_ids, start=1)
        }
        for student in context["students"]:
            student.row_number = row_numbers[student.pk]
        return context


class StudentPhoneFormMixin:
    def get_phone_formset(self, instance=None):
        instance = instance or getattr(self, "object", None)
        return StudentPhoneFormSet(
            data=self.request.POST if self.request.method == "POST" else None,
            instance=instance,
            prefix="phones",
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["phone_formset"] = kwargs.get("phone_formset") or self.get_phone_formset(
            instance=context["form"].instance
        )
        return context

    def form_invalid(self, form):
        phone_formset = self.get_phone_formset(instance=form.instance)
        return self.render_to_response(
            self.get_context_data(form=form, phone_formset=phone_formset)
        )

    def form_valid(self, form):
        phone_formset = self.get_phone_formset(instance=form.instance)
        if not phone_formset.is_valid():
            return self.form_invalid(form)

        student = form.save(commit=False)
        student.campus = get_selected_campus(self.request)
        student.phone = ", ".join(
            phone_form.cleaned_data["number"].strip()
            for phone_form in phone_formset.forms
            if phone_form.cleaned_data
            and not phone_form.cleaned_data.get("DELETE")
            and not phone_form.cleaned_data.get("invalid")
        )
        with transaction.atomic():
            student.save()
            phone_formset.instance = student
            phone_formset.save()
            form.save_m2m()
        self.object = student
        messages.success(self.request, self.success_message)
        return redirect(self.get_success_url())


class StudentCreateView(StudentPhoneFormMixin, CampusScopedMixin, CreateView):
    required_permission = "students.add_student"
    model = Student
    form_class = StudentForm
    template_name = "students/form.html"
    success_url = reverse_lazy("student-list")
    success_message = "Estudiante registrado correctamente."

    def get_initial(self):
        initial = super().get_initial()
        # Viene de «+ Nuevo estudiante» en la selección de matrícula.
        for field in ("last_name", "first_name"):
            value = self.request.GET.get(field, "").strip()[:120]
            if value:
                initial[field] = value.upper()
        return initial

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        kwargs["request"] = self.request
        return kwargs


class StudentUpdateView(StudentPhoneFormMixin, CampusScopedMixin, UpdateView):
    required_permission = "students.change_student"
    model = Student
    form_class = StudentForm
    template_name = "students/form.html"
    success_url = reverse_lazy("student-list")
    success_message = "Estudiante actualizado correctamente."

    def get_queryset(self):
        return self.campus_queryset(Student.objects.all())

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        kwargs["request"] = self.request
        return kwargs


class StudentDeleteView(CampusScopedMixin, DeleteView):
    required_permission = "core.deactivate_student"
    model = Student
    template_name = "students/student_confirm_delete.html"
    success_url = reverse_lazy("student-list")

    def get_queryset(self):
        return self.campus_queryset(Student.objects.all())

    def post(self, request, *args, **kwargs):
        student = self.get_object()
        self.object = student
        from payments.models import Installment, Payment

        with transaction.atomic():
            student.active = False
            student.save(update_fields=["active"])
            Enrollment.objects.filter(student=student, active=True).update(
                active=False, deactivated_with_student=True
            )
            Installment.objects.filter(enrollment__student=student).update(active=False)
            Payment.objects.filter(
                enrollment__student=student,
                status=Payment.Status.CONFIRMED,
            ).update(status=Payment.Status.VOID, voided_on_deactivation=True)
        messages.success(
            self.request,
            "Estudiante dado de baja; se conservaron sus matrículas y recibos, y los recibos confirmados se anularon.",
        )
        return redirect(self.get_success_url())


class EnrollmentListView(CampusScopedMixin, ListView):
    required_permission = "students.view_enrollment"
    model = Enrollment
    template_name = "students/enrollment_list.html"
    context_object_name = "enrollments"
    paginate_by = 25
    sortable_fields = {
        "row_number": "enrollment_date",
        "student": "student__full_name",
        "course": "course__name",
        "schedule": "course__schedule",
        "group": "course__group",
        "record_code": "ficha__code",
        "promoter": "promoter__name",
        "enrollment_date": "enrollment_date",
        "start_date": "course__start_date",
        "enrollment_fee": "enrollment_fee",
        "monthly_fee": "monthly_fee",
        "total_enrollment_cost": "total_enrollment_cost",
        "total_paid": "total_paid",
        "total_debt": "total_debt",
        "active": "active",
    }

    def get_queryset(self):
        from decimal import Decimal
        from payments.models import Installment, Payment, PaymentAllocation

        money_field = DecimalField(max_digits=12, decimal_places=2)
        confirmed_allocations = (
            PaymentAllocation.objects.filter(
                installment_id=OuterRef("pk"),
                payment__status=Payment.Status.CONFIRMED,
            )
            .values("installment_id")
            .annotate(total=Sum("amount"))
            .values("total")[:1]
        )
        installments = Installment.objects.annotate(
            paid_total=Coalesce(
                Subquery(confirmed_allocations, output_field=money_field),
                Value(Decimal("0.00"), output_field=money_field),
                output_field=money_field,
            )
        ).annotate(debt_balance=F("amount") - F("paid_total"))
        paid_totals = (
            installments.filter(enrollment_id=OuterRef("pk"))
            .order_by()
            .values("enrollment_id")
            .annotate(total=Sum("paid_total"))
            .values("total")[:1]
        )
        debt_totals = (
            installments.filter(enrollment_id=OuterRef("pk"))
            .order_by()
            .values("enrollment_id")
            .annotate(
                total=Sum(
                    Case(
                        When(debt_balance__gt=0, then=F("debt_balance")),
                        default=Value(Decimal("0.00"), output_field=money_field),
                        output_field=money_field,
                    )
                )
            )
            .values("total")[:1]
        )
        qs = self.campus_queryset(
            Enrollment.objects.select_related("campus", "student", "course", "ficha")
            .annotate(
                total_enrollment_cost=F("enrollment_fee") + F("monthly_fee"),
                total_paid=Coalesce(
                    Subquery(paid_totals, output_field=money_field),
                    Value(Decimal("0.00"), output_field=money_field),
                    output_field=money_field,
                ),
                total_debt=Coalesce(
                    Subquery(debt_totals, output_field=money_field),
                    Value(Decimal("0.00"), output_field=money_field),
                    output_field=money_field,
                ),
            )
            .prefetch_related(
                Prefetch("installments", queryset=installments),
                Prefetch("payments", queryset=Payment.objects.order_by("created_at", "pk")),
                "promoter",
            )
        )
        filters = {
            "student__full_name__unaccent__wild": self.request.GET.get("student", "").strip(),
            "course__name__unaccent__wild": self.request.GET.get("course", "").strip(),
            "ficha__code__icontains": self.request.GET.get("record_code", "").strip(),
            "promoter__name__unaccent__wild": self.request.GET.get("promoter", "").strip(),
        }
        for field, value in filters.items():
            if value:
                qs = qs.filter(**{field: value})
        for parameter, field in (
            ("course_id", "course_id"),
            ("promoter_id", "promoter__pk"),
            ("student_id", "student_id"),
        ):
            value = self.request.GET.get(parameter, "").strip()
            if value:
                try:
                    qs = qs.filter(**{field: int(value)})
                except ValueError:
                    qs = qs.none()
        enrollment_date = self.request.GET.get("enrollment_date", "").strip()
        if enrollment_date:
            parsed_date = parse_date(enrollment_date)
            qs = qs.filter(enrollment_date=parsed_date) if parsed_date else qs.none()
        for parameter, lookup in (
            ("enrollment_date_from", "enrollment_date__gte"),
            ("enrollment_date_to", "enrollment_date__lte"),
        ):
            value = self.request.GET.get(parameter, "").strip()
            if value:
                parsed_date = parse_date(value)
                qs = qs.filter(**{lookup: parsed_date}) if parsed_date else qs.none()
        debt_status = self.request.GET.get("debt_status", "")
        if debt_status == "with_debt":
            qs = qs.filter(total_debt__gt=0)
        elif debt_status == "without_debt":
            qs = qs.filter(total_debt=0)
        document = self.request.GET.get("document", "").strip()
        if document:
            qs = qs.filter(student__document_number__wild=document)
        schedule = self.request.GET.get("schedule", "").strip()
        if schedule:
            qs = qs.filter(course__schedule__unaccent__wild=schedule)
        group = self.request.GET.get("group", "").strip()
        if group:
            qs = qs.filter(course__group__iexact=group)
        record_status = self.request.GET.get("record_status", "")
        if record_status == "with_record":
            qs = qs.filter(ficha__isnull=False)
        elif record_status == "without_record":
            qs = qs.filter(ficha__isnull=True)
        elif record_status == "with_image":
            qs = qs.filter(ficha__isnull=False).exclude(ficha__image="")
        active_status = self.request.GET.get("active_status", "")
        if active_status == "active":
            qs = qs.filter(active=True)
        elif active_status == "inactive":
            qs = qs.filter(active=False)
        qs = qs.distinct()
        sort = self.request.GET.get("sort", "created_at")
        direction = self.request.GET.get("direction", "desc")
        sort_field = self.sortable_fields.get(sort, "created_at")
        if direction not in {"asc", "desc"}:
            direction = "desc"
        if sort == "row_number":
            chronological_fields = ["enrollment_date", "created_at", "pk"]
            if direction == "desc":
                chronological_fields = [f"-{field}" for field in chronological_fields]
            return qs.order_by(*chronological_fields)
        if direction == "desc":
            sort_field = f"-{sort_field}"
        return qs.order_by(sort_field, "pk")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        query_params = self.request.GET.copy()
        query_params.pop("page", None)
        context["pagination_query"] = query_params.urlencode()
        context["current_sort"] = self.request.GET.get("sort", "created_at")
        context["current_direction"] = self.request.GET.get("direction", "desc")
        context["course_filter_value"] = self.request.GET.get("course", "")
        course_id = self.request.GET.get("course_id", "")
        if course_id.isdigit():
            selected_course = Course.objects.filter(
                pk=course_id, campus=get_selected_campus(self.request)
            ).first()
            if selected_course:
                context["course_filter_value"] = selected_course.name
        context["promoter_filter_value"] = self.request.GET.get("promoter", "")
        promoter_id = self.request.GET.get("promoter_id", "")
        if promoter_id.isdigit():
            from .models import Promoter

            selected_promoter = Promoter.objects.filter(
                Q(campuses=get_selected_campus(self.request)) | Q(campuses__isnull=True),
                pk=promoter_id,
            ).first()
            if selected_promoter:
                context["promoter_filter_value"] = selected_promoter.name
        context["enrollment_date_filter_value"] = self.request.GET.get("enrollment_date", "")
        exact_enrollment_date = self.request.GET.get("enrollment_date", "")
        context["enrollment_date_from"] = (
            self.request.GET.get("enrollment_date_from", "") or exact_enrollment_date
        )
        context["enrollment_date_to"] = (
            self.request.GET.get("enrollment_date_to", "") or exact_enrollment_date
        )
        context["debt_filter_value"] = self.request.GET.get("debt_status", "")
        context["sort_urls"] = {}
        for key in self.sortable_fields:
            sort_params = self.request.GET.copy()
            sort_params.pop("page", None)
            if context["current_sort"] == key:
                next_direction = "desc" if context["current_direction"] == "asc" else "asc"
            else:
                next_direction = "asc"
            sort_params["sort"] = key
            sort_params["direction"] = next_direction
            context["sort_urls"][key] = f"?{sort_params.urlencode()}"
        chronological_ids = self.object_list.order_by(
            "enrollment_date", "created_at", "pk"
        ).values_list("pk", flat=True)
        row_numbers = {
            enrollment_id: row_number
            for row_number, enrollment_id in enumerate(chronological_ids, start=1)
        }
        for enrollment in context["enrollments"]:
            enrollment.row_number = row_numbers[enrollment.pk]
            enrollment.has_outstanding_debt = enrollment.active and enrollment.total_debt > 0
            enrollment.has_no_debt = enrollment.active and enrollment.total_debt == 0
        context["payable_enrollment_ids"] = set()
        if self.request.user.has_perm("payments.add_payment"):
            from payments.forms import pending_enrollments_for_campus

            context["payable_enrollment_ids"] = set(
                pending_enrollments_for_campus(get_selected_campus(self.request)).values_list(
                    "pk", flat=True
                )
            )
        return context


class ToolsView(
    FeaturePermissionMixin, SelectedCampusRequiredMixin, LoginRequiredMixin, TemplateView
):
    template_name = "utils/tools.html"

    def dispatch(self, request, *args, **kwargs):
        user = request.user
        if user.is_authenticated and not (
            user.has_perm("students.add_enrollment") or user.has_perm("payments.add_payment")
        ):
            from django.core.exceptions import PermissionDenied

            raise PermissionDenied("No tienes permiso para acceder a esta sección.")
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        context["can_check_record_codes"] = user.has_perm("students.add_enrollment")
        context["missing_record_codes"] = None
        context["record_code_check_error"] = None
        context["record_code_check_date"] = self.request.GET.get(
            "record_code_check_date", ""
        ).strip()
        context["record_code_start"] = self.request.GET.get("record_code_start", "").strip()
        context["record_code_end"] = self.request.GET.get("record_code_end", "").strip()
        if context["can_check_record_codes"] and self.request.GET.get("check_record_codes"):
            enrollment_date = parse_date(context["record_code_check_date"])
            if enrollment_date is None:
                context["record_code_check_error"] = "Seleccione una fecha de matrícula válida."
            elif not (
                context["record_code_start"].isdigit() and context["record_code_end"].isdigit()
            ):
                context["record_code_check_error"] = (
                    "Ingrese un rango numérico válido de códigos de ficha."
                )
            else:
                first_code, last_code = sorted(
                    (int(context["record_code_start"]), int(context["record_code_end"]))
                )
                if last_code - first_code + 1 > MAX_RECORD_CODE_CHECK_RANGE:
                    context["record_code_check_error"] = (
                        f"El rango no puede superar {MAX_RECORD_CODE_CHECK_RANGE:,} códigos."
                    )
                else:
                    expected_codes = list(range(first_code, last_code + 1))
                    registered_codes = set(
                        Ficha.objects.filter(
                            campus=get_selected_campus(self.request),
                            date=enrollment_date,
                            code__in=expected_codes,
                        ).values_list("code", flat=True)
                    )
                    context["missing_record_codes"] = [
                        code for code in expected_codes if code not in registered_codes
                    ]
        context["can_check_receipts"] = user.has_perm("payments.add_payment")
        context["missing_receipts"] = None
        context["receipt_check_error"] = None
        context["receipt_start"] = self.request.GET.get("receipt_start", "").strip()
        context["receipt_end"] = self.request.GET.get("receipt_end", "").strip()
        if context["can_check_receipts"] and self.request.GET.get("check_receipts"):
            from django.core.exceptions import ValidationError
            from payments.models import Payment
            from payments.views import MAX_RECEIPT_CHECK_RANGE, parse_receipt_code

            try:
                start_prefix, start_number = parse_receipt_code(context["receipt_start"])
                end_prefix, end_number = parse_receipt_code(context["receipt_end"])
                if start_prefix != end_prefix:
                    context["receipt_check_error"] = (
                        "Los dos recibos deben ser numéricos o usar la misma letra."
                    )
                else:
                    first_number, last_number = sorted((start_number, end_number))
                    if last_number - first_number + 1 > MAX_RECEIPT_CHECK_RANGE:
                        context["receipt_check_error"] = (
                            f"El rango no puede superar {MAX_RECEIPT_CHECK_RANGE:,} recibos."
                        )
                    else:
                        expected_receipts = [
                            f"{start_prefix}{number}"
                            for number in range(first_number, last_number + 1)
                        ]
                        registered_receipts = set(
                            Payment.objects.filter(
                                campus=get_selected_campus(self.request),
                                receipt_number__in=expected_receipts,
                            ).values_list("receipt_number", flat=True)
                        )
                        context["missing_receipts"] = [
                            receipt
                            for receipt in expected_receipts
                            if receipt not in registered_receipts
                        ]
            except (ValidationError, ValueError, TypeError):
                context["receipt_check_error"] = (
                    "Ingrese solo números o una letra mayúscula seguida de números."
                )
        return context


class StudentEnrollmentListView(EnrollmentListView):
    def get_student(self):
        if not hasattr(self, "student"):
            self.student = get_object_or_404(
                Student,
                pk=self.kwargs["student_id"],
                campus=get_selected_campus(self.request),
            )
        return self.student

    def get_queryset(self):
        self.get_student()
        return super().get_queryset().filter(student_id=self.student.pk)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["filtered_student"] = self.get_student()
        return context


class EnrollmentExcelExportView(EnrollmentListView):
    required_permission = ("students.view_enrollment", "core.export_reports")

    def get(self, request, *args, **kwargs):
        self.request = request
        enrollments = list(
            self.get_queryset().filter(active=True).order_by("student__full_name", "pk")
        )
        row_numbers = {
            enrollment.pk: row_number
            for row_number, enrollment in enumerate(
                enrollments,
                start=1,
            )
        }
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Matrículas"
        worksheet.sheet_view.showGridLines = False
        worksheet.merge_cells("A1:M1")
        worksheet["A1"] = "Reporte de matrículas"
        worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
        worksheet["A1"].fill = PatternFill("solid", fgColor="173A72")
        worksheet.merge_cells("A2:M2")
        worksheet["A2"] = (
            f"Sede: {get_selected_campus(request).name} · "
            f"Exportado: {timezone.localdate():%d/%m/%Y}"
        )
        headers = [
            "N.º",
            "Estudiante",
            "Curso",
            "Código de ficha",
            "Promotores",
            "Fecha de matrícula",
            "Inicio",
            "Inscripción",
            "Mensualidad",
            "Pago total",
            "Pagado",
            "Deuda",
            "Recibos",
        ]
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=4, column=column, value=header)
            cell.font = Font(bold=True, color="173A72")
            cell.fill = PatternFill("solid", fgColor="EAF1FF")

        for enrollment in enrollments:
            promoters = ", ".join(enrollment.promoter.values_list("name", flat=True)) or "—"
            receipts = (
                ", ".join(
                    f"{payment.receipt_number} (Anulado)"
                    if payment.status == "ANULADO"
                    else payment.receipt_number
                    for payment in enrollment.payments.all()
                )
                or "—"
            )
            worksheet.append(
                [
                    row_numbers[enrollment.pk],
                    enrollment.student_label,
                    enrollment.course_label,
                    enrollment.record_code or "—",
                    promoters,
                    enrollment.enrollment_date,
                    enrollment.course_start_date or EMPTY_VALUE,
                    enrollment.enrollment_fee,
                    enrollment.monthly_fee,
                    enrollment.total_enrollment_cost,
                    enrollment.total_paid,
                    enrollment.total_debt,
                    receipts,
                ]
            )
            if enrollment.total_debt > 0:
                fill_color = "FFF3CD"
            else:
                fill_color = "D1E7DD"
            for cell in worksheet[worksheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill_color)

        if enrollments:
            table = Table(
                displayName="EnrollmentTable",
                ref=f"A4:M{worksheet.max_row}",
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)
        else:
            worksheet["A4"] = "No hay matrículas que coincidan con los filtros."

        widths = [10, 30, 30, 18, 26, 18, 14, 15, 15, 15, 15, 15, 32]
        for column, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(column)].width = width
        worksheet.freeze_panes = "A5"
        for row in worksheet.iter_rows(min_row=5, min_col=6, max_col=7):
            for cell in row:
                cell.number_format = "DD/MM/YYYY"
        for row in worksheet.iter_rows(min_row=5, min_col=8, max_col=12):
            for cell in row:
                cell.number_format = '"S/" #,##0.00'

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="matriculas_{timezone.localdate():%Y%m%d}.xlsx"'
        )
        return response


class EnrollmentCreateView(CampusScopedMixin, CreateView):
    required_permission = "students.add_enrollment"
    model = Enrollment
    form_class = EnrollmentForm
    template_name = "students/form.html"
    selection_template_name = "students/enrollment_student_select.html"
    success_url = reverse_lazy("enrollment-list")
    students_per_page = 25

    def get_locked_student(self):
        if not hasattr(self, "_locked_student"):
            raw = self.request.POST.get("student") or self.request.GET.get("student")
            student = None
            if raw and raw.isdigit():
                student = Student.objects.filter(
                    pk=int(raw), campus=get_selected_campus(self.request), active=True
                ).first()
            self._locked_student = student
        return self._locked_student

    def is_without_student(self):
        return (
            self.request.POST.get("sin_estudiante") or self.request.GET.get("sin_estudiante")
        ) == "1"

    def get(self, request, *args, **kwargs):
        if self.get_locked_student() is None and not self.is_without_student():
            return self.render_student_selection()
        return super().get(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.POST.get("pending_course") == "1" and self.get_locked_student():
            # Matrícula rápida pendiente de asignación de curso.
            data = request.POST.copy()
            data.update(
                {
                    "course": "",
                    "enrollment_date": timezone.localdate().isoformat(),
                    "enrollment_fee": "0",
                    "monthly_fee": "0",
                    "active": "on",
                }
            )
            request.POST = data
        if self.get_locked_student() is None and not self.is_without_student():
            messages.error(request, "Primero selecciona un estudiante válido de la tabla.")
            return redirect("enrollment-create")
        return super().post(request, *args, **kwargs)

    def render_student_selection(self):
        students = Student.objects.filter(
            campus=get_selected_campus(self.request), active=True
        ).order_by("full_name", "pk")
        q = self.request.GET.get("q", "").strip()
        if q:
            students = filter_students(students, q)
        last_name = self.request.GET.get("apellidos", "").strip()
        first_name = self.request.GET.get("nombres", "").strip()
        for field, text in (("last_name", last_name), ("first_name", first_name)):
            for word in text.split()[:6]:
                # Los alumnos antiguos sin apellidos/nombres separados se buscan en el nombre completo.
                students = students.filter(
                    Q(**{f"{field}__unaccent__wild": word})
                    | Q(**{field: "", "full_name__unaccent__wild": word})
                )
        paginator = Paginator(students, self.students_per_page)
        page_obj = paginator.get_page(self.request.GET.get("page"))
        pagination_query = self.request.GET.copy()
        pagination_query.pop("page", None)
        return render(
            self.request,
            self.selection_template_name,
            {
                "students": page_obj.object_list,
                "page_obj": page_obj,
                "is_paginated": page_obj.has_other_pages(),
                "pagination_query": pagination_query.urlencode(),
                "q": q,
                "last_name": last_name,
                "first_name": first_name,
            },
        )

    def get_initial(self):
        initial = super().get_initial()
        initial["student"] = self.get_locked_student()
        return initial

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["locked_student"] = self.get_locked_student()
        context["without_student"] = self.is_without_student() and context["locked_student"] is None
        if context["locked_student"]:
            context["current_enrollments"] = (
                Enrollment.objects.filter(
                    student=context["locked_student"], campus=get_selected_campus(self.request)
                )
                .select_related("course")
                .order_by("-enrollment_date", "pk")
            )
        return context

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        kwargs["request"] = self.request
        kwargs["locked_student"] = self.get_locked_student()
        kwargs["without_student"] = self.is_without_student() and kwargs["locked_student"] is None
        return kwargs

    def form_valid(self, form):
        campus = get_selected_campus(self.request)
        form.instance.campus = campus
        form.instance.enrollment_fee = form.cleaned_data["enrollment_fee"]
        form.instance.monthly_fee = form.cleaned_data["monthly_fee"]
        try:
            with transaction.atomic():
                self.object = form.save()
                from payments.models import Installment

                due_date = self.object.course_start_date or self.object.enrollment_date
                if self.object.enrollment_fee > 0:
                    Installment.objects.create(
                        enrollment=self.object,
                        concept="Matrícula",
                        amount=self.object.enrollment_fee,
                        due_date=due_date,
                    )
                if self.object.monthly_fee > 0:
                    Installment.objects.create(
                        enrollment=self.object,
                        concept=f"Mensualidad inicial ({due_date:%Y-%m})",
                        amount=self.object.monthly_fee,
                        due_date=due_date,
                    )
        except IntegrityError:
            form.add_error("course", "Este estudiante ya está matriculado en este curso.")
            return self.form_invalid(form)
        messages.success(self.request, "Matrícula registrada y cuotas iniciales generadas.")
        message = seats_message(self.object.course)
        if message:
            messages.info(self.request, message)
        return redirect(self.get_success_url())


class EnrollmentUpdateView(CampusScopedMixin, UpdateView):
    required_permission = "students.change_enrollment"
    model = Enrollment
    form_class = EnrollmentForm
    template_name = "students/form.html"
    success_url = reverse_lazy("enrollment-list")

    def get_required_permissions(self, request):
        # Quien puede matricular también completa las matrículas pendientes.
        user = request.user
        if user.is_authenticated and user.has_perm("students.add_enrollment"):
            pk = self.kwargs.get("pk")
            enrollment = (
                Enrollment.objects.select_related("course").filter(pk=pk).first() if pk else None
            )
            if enrollment and enrollment.is_pending:
                return ("students.add_enrollment",)
        return super().get_required_permissions(request)

    def get_queryset(self):
        return self.campus_queryset(
            Enrollment.objects.select_related("campus", "student", "course").prefetch_related(
                "promoter"
            )
        )

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        kwargs["request"] = self.request
        return kwargs

    def form_valid(self, form):
        form.instance.campus = get_selected_campus(self.request)
        form.instance.enrollment_fee = form.cleaned_data["enrollment_fee"]
        form.instance.monthly_fee = form.cleaned_data["monthly_fee"]
        was_active = (
            Enrollment.objects.filter(pk=form.instance.pk).values_list("active", flat=True).get()
        )
        from payments.models import Installment

        with transaction.atomic():
            # Bloquea las cuotas para validar contra pagos concurrentes.
            installments = list(
                Installment.objects.select_for_update().filter(enrollment=form.instance)
            )
            fee_errors = self.fee_conflicts(form, installments)
            if fee_errors:
                for field, message in fee_errors:
                    form.add_error(field, message)
                return self.form_invalid(form)
            response = super().form_valid(form)
            if not was_active and self.object.active:
                Installment.objects.filter(enrollment=self.object).update(active=True)
            self.sync_installments(installments)
        messages.success(self.request, "Matrícula actualizada y deuda recalculada.")
        message = seats_message(self.object.course)
        if message:
            messages.info(self.request, message)
        return response

    def fee_conflicts(self, form, installments):
        return fee_conflicts(installments, form.cleaned_data)

    def sync_installments(self, installments):
        sync_initial_installments(self.object, installments)


class EnrollmentDeleteView(CampusScopedMixin, DeleteView):
    required_permission = "core.deactivate_enrollment"
    model = Enrollment
    template_name = "students/enrollment_confirm_delete.html"
    success_url = reverse_lazy("enrollment-list")

    def get_queryset(self):
        return self.campus_queryset(
            Enrollment.objects.select_related("campus", "student", "course")
        )

    def post(self, request, *args, **kwargs):
        enrollment = self.get_object()
        self.object = enrollment
        from payments.models import Installment, Payment

        with transaction.atomic():
            enrollment.active = False
            enrollment.save(update_fields=["active"])
            Installment.objects.filter(enrollment=enrollment).update(active=False)
            Payment.objects.filter(
                enrollment=enrollment,
                status=Payment.Status.CONFIRMED,
            ).update(status=Payment.Status.VOID, voided_on_deactivation=True)
        messages.success(
            self.request,
            "Matrícula dada de baja; se conservaron las cuotas y recibos, y los recibos confirmados se anularon.",
        )
        return redirect(self.get_success_url())


def seats_message(course):
    if not course or not course.capacity:
        return None
    occupied = (
        Enrollment.objects.filter(course=course, active=True)
        .filter(Q(student__isnull=True) | Q(student__active=True))
        .count()
    )
    left = course.capacity - occupied
    label = f"{course.name} (grupo {course.group})" if course.group else course.name
    if left > 0:
        noun = "vacante" if left == 1 else "vacantes"
        return f"Quedan {left} {noun} disponibles en {label}."
    if left == 0:
        return f"El curso {label} ya no tiene vacantes disponibles."
    return f"El curso {label} excede su capacidad en {-left}."


def course_students_queryset(request):
    campus = get_selected_campus(request)
    enrollments = (
        Enrollment.objects.filter(campus=campus, active=True)
        .filter(Q(student__isnull=True) | Q(student__active=True))
        .select_related("student", "course")
        .order_by("student__full_name")
    )
    qs = (
        Course.objects.filter(campus=campus, active=True)
        .annotate(
            student_count=Count(
                "enrollments",
                filter=Q(enrollments__active=True, enrollments__student__active=True),
            ),
            enrollment_count=Count(
                "enrollments",
                filter=Q(enrollments__active=True)
                & (Q(enrollments__student__isnull=True) | Q(enrollments__student__active=True)),
            ),
            first_enrollment_date=Min(
                "enrollments__enrollment_date", filter=Q(enrollments__active=True)
            ),
            last_enrollment_date=Max(
                "enrollments__enrollment_date", filter=Q(enrollments__active=True)
            ),
        )
        .prefetch_related(Prefetch("enrollments", queryset=enrollments))
        .order_by("name", "schedule", "group")
    )
    query = request.GET.get("q", "").strip()
    if query:
        qs = qs.filter(
            Q(name__unaccent__wild=query)
            | Q(schedule__unaccent__wild=query)
            | Q(group__unaccent__wild=query)
        )
    return qs


def course_seat_summary(courses):
    summary = []
    for course in courses:
        capacity = course.capacity
        occupied = course.enrollment_count
        summary.append(
            {
                "course": course,
                "occupied": occupied,
                "reserved": occupied - course.student_count,
                "available": max(capacity - occupied, 0) if capacity else None,
                "over": bool(capacity) and occupied > capacity,
                "occupancy_percent": occupied * 100 / capacity if capacity else None,
                "first_enrollment_date": course.first_enrollment_date,
                "last_enrollment_date": course.last_enrollment_date,
            }
        )
    return summary


class CourseStudentTablesView(CampusScopedMixin, ListView):
    required_permission = "core.view_academic_reports"
    """Muestra una tabla independiente por curso para la sede seleccionada."""

    model = Course
    template_name = "students/course_student_tables.html"
    context_object_name = "courses"
    paginate_by = 10

    def get_queryset(self):
        return course_students_queryset(self.request)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["search_query"] = self.request.GET.get("q", "").strip()
        params = self.request.GET.copy()
        params.pop("page", None)
        context["pagination_query"] = params.urlencode()
        context["seat_summary"] = course_seat_summary(course_students_queryset(self.request))
        context["total_students"] = (
            Enrollment.objects.filter(
                campus=get_selected_campus(self.request), active=True, student__active=True
            )
            .values("student_id")
            .distinct()
            .count()
        )
        return context


class CourseStudentExportView(CampusScopedMixin, View):
    required_permission = ("core.view_academic_reports", "core.export_reports")

    def get(self, request, *args, **kwargs):
        workbook = Workbook()
        all_courses = list(course_students_queryset(request))
        courses = [course for course in all_courses if course.enrollment_count > 0]
        seat_summary = course_seat_summary(all_courses)
        grouped_courses = [
            (name, list(course_group))
            for name, course_group in groupby(courses, key=lambda course: course.name)
        ]
        workbook.remove(workbook.active)
        campus_name = get_selected_campus(request).name
        export_date = timezone.localdate()
        used_titles = {"ocupación"}

        occupancy_sheet = workbook.create_sheet("Ocupación")
        occupancy_sheet.sheet_view.showGridLines = False
        occupancy_sheet.merge_cells("A1:H1")
        occupancy_sheet["A1"] = "Ocupación y vacantes por curso"
        occupancy_sheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
        occupancy_sheet["A1"].fill = PatternFill("solid", fgColor="173A72")
        occupancy_sheet.merge_cells("A2:H2")
        occupancy_sheet["A2"] = f"Sede: {campus_name} · Exportado: {export_date:%d/%m/%Y}"
        occupancy_headers = [
            "Curso",
            "Grupo",
            "Horario",
            "Capacidad",
            "Ocupadas",
            "Ocupación",
            "Primera matrícula",
            "Última matrícula",
        ]
        for column, header in enumerate(occupancy_headers, start=1):
            cell = occupancy_sheet.cell(row=4, column=column, value=header)
            cell.font = Font(bold=True, color="173A72")
            cell.fill = PatternFill("solid", fgColor="EAF1FF")
        for row_number, item in enumerate(seat_summary, start=5):
            course = item["course"]
            row = [
                course.name,
                course.group or EMPTY_VALUE,
                course.schedule,
                course.capacity or EMPTY_VALUE,
                item["occupied"],
                item["occupancy_percent"] / 100 if item["occupancy_percent"] is not None else None,
                item["first_enrollment_date"],
                item["last_enrollment_date"],
            ]
            for column, value in enumerate(row, start=1):
                occupancy_sheet.cell(row=row_number, column=column, value=value)
            occupancy_sheet.cell(row=row_number, column=6).number_format = "0%"
            for column in (7, 8):
                occupancy_sheet.cell(row=row_number, column=column).number_format = "DD/MM/YYYY"
        occupancy_sheet.freeze_panes = "A5"
        for column, width in enumerate((32, 16, 24, 14, 14, 14, 22, 22), start=1):
            occupancy_sheet.column_dimensions[get_column_letter(column)].width = width

        if not grouped_courses:
            worksheet = workbook.create_sheet("Sin cursos")
            worksheet["A1"] = "No hay cursos activos para exportar."

        for course_name, same_name_courses in grouped_courses:
            base_title = re.sub(r"[\[\]:*?/\\]", " ", course_name)
            base_title = (
                "".join(
                    character if character.isprintable() else " " for character in base_title
                ).strip(" '")
                or "Curso"
            )
            title = base_title[:31]
            suffix_number = 2
            while title.casefold() in used_titles:
                suffix = f" ({suffix_number})"
                title = f"{base_title[: 31 - len(suffix)]}{suffix}"
                suffix_number += 1
            used_titles.add(title.casefold())

            worksheet = workbook.create_sheet(title)
            worksheet.sheet_view.showGridLines = False
            worksheet.merge_cells("A1:E1")
            worksheet["A1"] = course_name
            worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
            worksheet["A1"].fill = PatternFill("solid", fgColor="173A72")
            worksheet.merge_cells("A2:E2")
            worksheet["A2"] = f"Sede: {campus_name} · Exportado: {export_date:%d/%m/%Y}"

            for course in same_name_courses:
                section_row = worksheet.max_row + 2
                section_details = [
                    f"Grupo: {course.group or '—'}",
                    f"Horario: {course.schedule or '—'}",
                ]
                section_details.extend(
                    [
                        f"Capacidad: {course.capacity_display()}",
                        f"Inicio: {course.start_date_display()}",
                    ]
                )
                worksheet.merge_cells(
                    start_row=section_row, start_column=1, end_row=section_row, end_column=5
                )
                section = worksheet.cell(
                    row=section_row, column=1, value=" · ".join(section_details)
                )
                section.font = Font(bold=True, color="FFFFFF")
                section.fill = PatternFill("solid", fgColor="1E9E6A")

                headers = ["N.º", "Alumno", "DNI", "Teléfonos", "Fecha de matrícula"]
                header_row = section_row + 1
                for column, header in enumerate(headers, start=1):
                    cell = worksheet.cell(row=header_row, column=column, value=header)
                    cell.font = Font(bold=True, color="173A72")
                    cell.fill = PatternFill("solid", fgColor="EAF1FF")

                course_enrollments = list(course.enrollments.all())
                for row_number, enrollment in enumerate(course_enrollments, start=1):
                    row = [
                        row_number,
                        enrollment.student.full_name
                        if enrollment.student_id
                        else "VACANTE RESERVADA",
                        enrollment.student.document_number
                        if enrollment.student_id
                        else EMPTY_VALUE,
                        enrollment.student.phone if enrollment.student_id else EMPTY_VALUE,
                        enrollment.enrollment_date,
                    ]
                    worksheet.append(row)
                    if not enrollment.student_id:
                        for cell in worksheet[worksheet.max_row]:
                            cell.fill = PatternFill("solid", fgColor="FFF3CD")

                table = Table(
                    displayName=f"Course_{course.pk}",
                    ref=f"A{header_row}:E{worksheet.max_row}",
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                worksheet.add_table(table)

            worksheet.freeze_panes = "A5"
            widths = [10, 32, 18, 16, 22]
            for column, width in enumerate(widths, start=1):
                worksheet.column_dimensions[get_column_letter(column)].width = width
            for row in worksheet.iter_rows(min_row=5):
                row[4].number_format = "DD/MM/YYYY"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="alumnos_por_curso_{timezone.localdate():%Y%m%d}.xlsx"'
        )
        workbook.save(response)
        return response


class AvailableStudentsForCourseView(CampusScopedMixin, ListView):
    required_permission = "students.add_enrollment"
    """Devuelve por AJAX los alumnos de la sede aún no inscritos al curso."""

    def get(self, request, *args, **kwargs):
        campus = get_selected_campus(request)
        course_id = request.GET.get("course_id")
        if not course_id:
            return JsonResponse({"students": []})
        course = Course.objects.filter(pk=course_id, campus=campus, active=True).first()
        if course is None:
            return JsonResponse({"students": []}, status=400)
        enrollments = Enrollment.objects.filter(course=course)
        current_enrollment = None
        enrollment_id = request.GET.get("enrollment_id")
        if enrollment_id:
            current_enrollment = enrollments.filter(pk=enrollment_id, campus=campus).first()
        if current_enrollment:
            enrollments = enrollments.exclude(pk=current_enrollment.pk)
        enrolled_ids = enrollments.filter(student__isnull=False).values_list(
            "student_id", flat=True
        )
        students = Student.objects.filter(campus=campus, active=True)
        if current_enrollment:
            students = Student.objects.filter(campus=campus).filter(
                Q(active=True) | Q(pk=current_enrollment.student_id)
            )
        students = (
            students.exclude(pk__in=enrolled_ids).order_by("full_name").values("id", "full_name")
        )
        return JsonResponse({"students": list(students)})
