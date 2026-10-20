from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.shortcuts import redirect
from django.views.generic import TemplateView
from openpyxl import load_workbook

from core.activity import record
from core.campus import FeaturePermissionMixin, SelectedCampusRequiredMixin, get_selected_campus
from core.models import ActivityLog

from .duplicates import name_key
from .models import Course, Enrollment, Student, create_student_phones_from_legacy

SESSION_KEY = "student_import_rows"
MAX_ROWS = 2000
HEADERS = {
    "nombre": "name",
    "nombres": "name",
    "nombre completo": "name",
    "alumno": "name",
    "estudiante": "name",
    "dni": "document",
    "documento": "document",
    "telefono": "phone",
    "celular": "phone",
    "curso": "course",
}


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _header_key(value):
    return name_key(_clean(value)).lower() if value else ""


def read_rows(file):
    """Lee la primera hoja y devuelve filas {name, document, phone, course} según los encabezados."""
    sheet = load_workbook(file, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    columns = {
        HEADERS[_header_key(h)]: i for i, h in enumerate(header) if _header_key(h) in HEADERS
    }
    if "name" not in columns:
        raise ValueError("Falta la columna «Nombre» en la primera fila.")
    parsed = []
    for row in rows:
        item = {key: _clean(row[i]) if i < len(row) else "" for key, i in columns.items()}
        if any(item.values()):
            parsed.append({"name": "", "document": "", "phone": "", "course": "", **item})
        if len(parsed) > MAX_ROWS:
            raise ValueError(f"El archivo supera el máximo de {MAX_ROWS} filas.")
    return parsed


def classify(rows, campus):
    """Marca cada fila como nueva, duplicada o con error, sin guardar nada."""
    existing = {
        name_key(n)
        for n in Student.objects.filter(campus=campus).values_list("full_name", flat=True)
    }
    courses = {}
    for course in Course.objects.filter(campus=campus):
        courses.setdefault(course.name.strip().lower(), []).append(course)
    seen = set()
    result = []
    for number, row in enumerate(rows, start=2):
        row = {**row, "line": number, "state": "new", "message": "", "course_id": None}
        key = name_key(row["name"])
        if not key:
            row.update(state="error", message="Nombre vacío.")
        elif key in existing or key in seen:
            row.update(state="duplicate", message="Ya existe un estudiante con ese nombre.")
        else:
            student = Student(
                campus=campus,
                full_name=row["name"],
                phone=row["phone"],
                document_number=row["document"],
            )
            try:
                student.full_clean(exclude=["campus"], validate_unique=False)
            except ValidationError as exc:
                row.update(
                    state="error",
                    message=" ".join(m for msgs in exc.message_dict.values() for m in msgs),
                )
            if row["state"] == "new" and row["course"]:
                matches = courses.get(row["course"].lower(), [])
                if len(matches) == 1:
                    row["course_id"] = matches[0].pk
                else:
                    row.update(
                        state="error",
                        message="Curso inexistente."
                        if not matches
                        else "Hay varios cursos con ese nombre.",
                    )
        if row["state"] == "new":
            seen.add(key)
        result.append(row)
    return result


class StudentImportView(FeaturePermissionMixin, SelectedCampusRequiredMixin, TemplateView):
    """Importa estudiantes (y opcionalmente su matrícula) desde Excel con vista previa."""

    required_permission = "students.add_student"
    template_name = "students/import.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        rows = kwargs.get("rows")
        context["rows"] = rows
        if rows is not None:
            context["counts"] = {
                state: sum(1 for r in rows if r["state"] == state)
                for state in ("new", "duplicate", "error")
            }
        return context

    def post(self, request, *args, **kwargs):
        campus = get_selected_campus(request)
        if request.POST.get("accion") == "confirmar":
            return self.confirm(request, campus)
        upload = request.FILES.get("archivo")
        if not upload or not upload.name.lower().endswith(".xlsx"):
            messages.error(request, "Sube un archivo Excel (.xlsx).")
            return redirect("student-import")
        try:
            rows = classify(read_rows(upload), campus)
        except ValueError as exc:
            messages.error(request, f"No se pudo leer el archivo: {exc}")
            return redirect("student-import")
        except Exception:
            messages.error(request, "No se pudo leer el archivo: no es un Excel válido.")
            return redirect("student-import")
        request.session[SESSION_KEY] = [r for r in rows if r["state"] == "new"]
        return self.render_to_response(self.get_context_data(rows=rows))

    def confirm(self, request, campus):
        rows = request.session.pop(SESSION_KEY, [])
        if not rows:
            messages.error(request, "No hay una vista previa pendiente. Sube el archivo de nuevo.")
            return redirect("student-import")
        can_enroll = request.user.has_perm("students.add_enrollment")
        created = enrolled = skipped = 0
        for row in rows:
            try:
                with transaction.atomic():
                    student = Student.objects.create(
                        campus=campus,
                        full_name=row["name"],
                        phone=row["phone"],
                        document_number=row["document"],
                    )
                    create_student_phones_from_legacy(student)
                    created += 1
                    course = Course.objects.filter(pk=row["course_id"], campus=campus).first()
                    if course and can_enroll:
                        enrollment = Enrollment.objects.create(
                            campus=campus,
                            student=student,
                            course=course,
                            enrollment_fee=course.enrollment_fee,
                            monthly_fee=course.monthly_fee,
                        )
                        self.create_installments(enrollment)
                        enrolled += 1
            except IntegrityError:
                skipped += 1
        text = f"Importación lista: {created} estudiante(s) y {enrolled} matrícula(s)."
        if skipped:
            text += f" {skipped} fila(s) omitidas por duplicado."
        record(
            request,
            ActivityLog.Action.CREATE,
            f"Importó desde Excel {created} estudiante(s) y {enrolled} matrícula(s)",
        )
        messages.success(request, text)
        return redirect("student-list")

    @staticmethod
    def create_installments(enrollment):
        from payments.models import Installment

        due_date = enrollment.course_start_date or enrollment.enrollment_date
        if enrollment.enrollment_fee > 0:
            Installment.objects.create(
                enrollment=enrollment,
                concept="Matrícula",
                amount=enrollment.enrollment_fee,
                due_date=due_date,
            )
        if enrollment.monthly_fee > 0:
            Installment.objects.create(
                enrollment=enrollment,
                concept=f"Mensualidad inicial ({due_date:%Y-%m})",
                amount=enrollment.monthly_fee,
                due_date=due_date,
            )
