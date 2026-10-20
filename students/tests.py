from datetime import date
from io import BytesIO

import threading

from django.db import IntegrityError, connections, transaction
from django.core.exceptions import ValidationError
from django.test import TestCase, TransactionTestCase, skipUnlessDBFeature
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Campus, User
from core.testing import create_user, grant, revoke
from payments.models import Installment, Payment, PaymentAllocation
from students.models import (
    Course,
    Enrollment,
    Ficha,
    Promoter,
    Schedule,
    Student,
    StudentPhone,
    create_student_phones_from_legacy,
)


class CourseStudentExportTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede Exportación", active=True)
        self.user = create_user(
            username="exporter",
            password="secret123",
            manage_students=True,
        )
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()

        self.student_a = Student.objects.create(
            campus=self.campus,
            full_name="Ana Pérez",
            document_number="12345678",
            phone="987654321",
        )
        self.student_b = Student.objects.create(
            campus=self.campus,
            full_name="Luis Gómez",
            document_number="87654321",
            phone="912345678",
        )
        self.student_c = Student.objects.create(
            campus=self.campus,
            full_name="Marta Ruiz",
            document_number="23456789",
            phone="923456789",
        )
        self.course_a = Course.objects.create(
            campus=self.campus,
            name="Matemáticas",
            schedule="Lunes",
            group="A",
            start_date=date(2026, 10, 1),
        )
        course_b = Course.objects.create(
            campus=self.campus,
            name="Matemáticas",
            schedule="Miércoles",
            group="B",
            start_date=date(2026, 10, 1),
        )
        course_c = Course.objects.create(
            campus=self.campus,
            name="Matemáticas",
            schedule="Martes",
            group="C",
            start_date=date(2026, 10, 1),
        )
        course_d = Course.objects.create(
            campus=self.campus,
            name="Matemáticas",
            schedule="Jueves",
            group="D",
            start_date=date(2026, 10, 1),
        )
        course_e = Course.objects.create(
            campus=self.campus,
            name="Matemáticas",
            schedule="Viernes",
            group="E",
            start_date=date(2026, 10, 1),
        )
        course_f = Course.objects.create(
            campus=self.campus,
            name="Historia",
            schedule="Viernes",
            group="F",
            start_date=date(2026, 10, 1),
        )
        Course.objects.create(
            campus=self.campus,
            name="Física",
            schedule="Sábado",
            group="G",
            start_date=date(2026, 10, 1),
        )
        Enrollment.objects.create(campus=self.campus, student=self.student_a, course=self.course_a)
        Enrollment.objects.create(campus=self.campus, student=self.student_b, course=self.course_a)
        Enrollment.objects.create(campus=self.campus, student=self.student_b, course=course_b)
        Enrollment.objects.create(campus=self.campus, student=self.student_c, course=course_c)
        Enrollment.objects.create(campus=self.campus, student=self.student_a, course=course_d)
        Enrollment.objects.create(campus=self.campus, student=self.student_b, course=course_e)
        Enrollment.objects.create(campus=self.campus, student=self.student_c, course=course_f)

    def test_export_creates_one_sheet_per_course_name(self):
        response = self.client.get(reverse("course-student-export"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Ocupación", "HISTORIA", "MATEMÁTICAS"])
        occupancy = list(workbook["Ocupación"].iter_rows(values_only=True))
        self.assertEqual(
            occupancy[3],
            (
                "Curso",
                "Grupo",
                "Horario",
                "Capacidad",
                "Ocupadas",
                "Ocupación",
                "Primera matrícula",
                "Última matrícula",
            ),
        )
        algebra_summary = next(
            row for row in occupancy[4:] if row[0] == "MATEMÁTICAS" and row[1] == "A"
        )
        self.assertEqual(algebra_summary[4:6], (2, 0.1))
        self.assertNotIn("FÍSICA", workbook.sheetnames)
        math_rows = list(workbook["MATEMÁTICAS"].iter_rows(values_only=True))
        history_rows = list(workbook["HISTORIA"].iter_rows(values_only=True))
        self.assertEqual(len(workbook["MATEMÁTICAS"].tables), 5)
        self.assertEqual(len(workbook["HISTORIA"].tables), 1)
        self.assertIn("ANA PÉREZ", [value for row in math_rows for value in row])
        self.assertIn("LUIS GÓMEZ", [value for row in math_rows for value in row])
        math_values = [value for row in math_rows for value in row]
        for group in ("A", "B", "C", "D", "E"):
            self.assertTrue(
                any(
                    isinstance(value, str) and value.startswith(f"Grupo: {group} ·")
                    for value in math_values
                )
            )
        self.assertIn("MARTA RUIZ", [value for row in history_rows for value in row])
        math_table = workbook["MATEMÁTICAS"].tables[f"Course_{self.course_a.pk}"]
        header_row = int(math_table.ref.split(":")[0][1:])
        self.assertEqual(
            [workbook["MATEMÁTICAS"].cell(header_row, column).value for column in range(1, 6)],
            ["N.º", "Alumno", "DNI", "Teléfonos", "Fecha de matrícula"],
        )
        self.assertEqual(workbook["MATEMÁTICAS"].cell(header_row + 1, 1).value, 1)
        self.assertEqual(workbook["MATEMÁTICAS"].cell(header_row + 2, 1).value, 2)

    def test_course_student_report_shows_reserved_seats_as_distinct_rows(self):
        course = Course.objects.get(campus=self.campus, name="FÍSICA")
        Enrollment.objects.create(
            campus=self.campus, course=course, enrollment_date=date(2026, 9, 10)
        )

        response = self.client.get(reverse("course-student-tables"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "FÍSICA")
        self.assertContains(response, "Vacante reservada · Sin estudiante")
        course_in_context = next(
            item for item in response.context["courses"] if item.pk == course.pk
        )
        self.assertEqual(course_in_context.enrollment_count, 1)
        self.assertEqual(course_in_context.student_count, 0)

    def test_course_student_report_has_vacancy_summary(self):
        course = Course.objects.get(campus=self.campus, name="FÍSICA")
        course.capacity = 3
        course.save()
        Enrollment.objects.create(
            campus=self.campus, course=course, enrollment_date=date(2026, 9, 10)
        )

        response = self.client.get(reverse("course-student-tables"))

        self.assertContains(response, "Ocupación y vacantes por curso")
        row = next(r for r in response.context["seat_summary"] if r["course"].pk == course.pk)
        self.assertEqual(row["occupied"], 1)
        self.assertEqual(row["reserved"], 1)
        self.assertEqual(row["available"], 2)
        self.assertFalse(row["over"])
        self.assertAlmostEqual(row["occupancy_percent"], 100 / 3)
        self.assertEqual(row["first_enrollment_date"], date(2026, 9, 10))
        self.assertEqual(row["last_enrollment_date"], date(2026, 9, 10))
        self.assertContains(response, "Primera matrícula")
        self.assertContains(response, "Ocupación")

    def test_course_student_export_includes_and_highlights_reserved_seats(self):
        course = Course.objects.get(campus=self.campus, name="FÍSICA")
        Enrollment.objects.create(campus=self.campus, course=course)

        response = self.client.get(reverse("course-student-export"))

        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook["FÍSICA"]
        rows = list(worksheet.iter_rows(values_only=True))
        self.assertIn("VACANTE RESERVADA", [value for row in rows for value in row])
        reserved_row = next(
            row for row in worksheet.iter_rows() if row[1].value == "VACANTE RESERVADA"
        )
        self.assertTrue(reserved_row[1].fill.fgColor.rgb.endswith("FFF3CD"))


class ExtraReportsTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede Informes", active=True)
        self.user = create_user(
            username="reporter",
            password="x-Test-pass-123",
            manage_students=True,
            register_payments=True,
        )
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.course = Course.objects.create(
            campus=self.campus, name="Álgebra", schedule="Lunes", start_date=date(2026, 10, 1)
        )
        self.promoter = Promoter.objects.create(name="Promotor Uno")
        self.promoter.campuses.add(self.campus)
        self.student = Student.objects.create(
            campus=self.campus, full_name="Ana Pérez", document_number="1", phone="987654321"
        )
        self.with_course = Enrollment.objects.create(
            campus=self.campus, student=self.student, course=self.course
        )
        self.with_course.promoter.add(self.promoter)
        self.pending = Enrollment.objects.create(
            campus=self.campus,
            student=Student.objects.create(
                campus=self.campus, full_name="Luis Gómez", document_number="2", phone="912345678"
            ),
        )
        self.seat = Enrollment.objects.create(campus=self.campus, course=self.course)
        Payment.objects.create(
            campus=self.campus,
            enrollment=self.with_course,
            cashier=self.user,
            receipt_number="R-1",
            amount=50,
        )
        Payment.objects.create(
            campus=self.campus,
            enrollment=self.with_course,
            cashier=self.user,
            receipt_number="R-2",
            amount=30,
            status=Payment.Status.VOID,
        )

    def test_reports_render_with_expected_figures(self):
        response = self.client.get(reverse("report-enrollments-by-promoter"))
        self.assertEqual(response.status_code, 200)
        self.assertIn(["PROMOTOR UNO", 1, 1, 0], response.context["rows"])
        self.assertIn(["Sin promotor", 2, 1, 1], response.context["rows"])

        response = self.client.get(reverse("report-enrollments-by-month"))
        self.assertEqual(response.context["totals"], ["Total", 3, 2, 1, 1])

        response = self.client.get(reverse("report-pending-course"))
        self.assertEqual([row[0] for row in response.context["rows"]], ["LUIS GÓMEZ"])

        response = self.client.get(reverse("report-income-by-month"))
        self.assertEqual(response.context["totals"], ["Total", 1, "50.00"])

        response = self.client.get(reverse("report-course-occupancy"))
        self.assertRedirects(response, reverse("course-student-tables"))
        response = self.client.get(reverse("course-student-tables"))
        row = next(
            item for item in response.context["seat_summary"] if item["course"].name == "ÁLGEBRA"
        )
        self.assertEqual(row["occupied"], 2)
        self.assertEqual(row["occupancy_percent"], 10)
        self.assertRedirects(
            self.client.get(f"{reverse('report-course-occupancy')}?export=xlsx"),
            reverse("course-student-export"),
        )

        Installment.objects.create(
            enrollment=self.with_course, concept="Cuota", amount=80, due_date=date(2026, 10, 1)
        )
        response = self.client.get(reverse("report-debt-by-course"))
        self.assertEqual(response.context["totals"], ["Total", "", "", 1, "80.00"])

    def test_financial_report_requires_permission(self):
        limited = create_user(username="limited", password="x-Test-pass-123", manage_students=True)
        limited.campuses.add(self.campus)
        self.client.force_login(limited)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.assertEqual(self.client.get(reverse("report-income-by-month")).status_code, 403)
        self.assertEqual(self.client.get(reverse("report-pending-course")).status_code, 200)


class LifecycleTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede Ciclo", active=True)
        self.user = create_user(
            username="lifecycle",
            password="x-Test-pass-123",
            manage_students=True,
            register_payments=True,
        )
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.course_a = Course.objects.create(
            campus=self.campus,
            name="Curso A",
            schedule="Lunes",
            start_date=date(2026, 10, 1),
            enrollment_fee=40,
            monthly_fee=60,
        )
        self.course_b = Course.objects.create(
            campus=self.campus,
            name="Curso B",
            schedule="Martes",
            start_date=date(2026, 11, 2),
            enrollment_fee=30,
            monthly_fee=70,
        )
        self.student = Student.objects.create(
            campus=self.campus, full_name="Ana Pérez", document_number="1", phone="987654321"
        )
        self.enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=self.student,
            course=self.course_a,
            enrollment_fee=40,
            monthly_fee=60,
        )
        self.fee = Installment.objects.create(
            enrollment=self.enrollment, concept="Matrícula", amount=40, due_date=date(2026, 10, 1)
        )
        self.monthly = Installment.objects.create(
            enrollment=self.enrollment,
            concept="Mensualidad inicial (2026-10)",
            amount=60,
            due_date=date(2026, 10, 1),
        )
        self.payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="R-1",
            amount=40,
        )
        PaymentAllocation.objects.create(payment=self.payment, installment=self.fee, amount=40)

    def transfer(self, course, use_fees):
        data = {"course": course.pk}
        if use_fees:
            data["use_course_fees"] = "on"
        return self.client.post(reverse("enrollment-transfer", args=[self.enrollment.pk]), data)

    def test_transfer_is_blocked_when_new_fee_is_below_paid_amount(self):
        response = self.transfer(self.course_b, use_fees=True)
        self.assertEqual(response.status_code, 200)
        self.enrollment.refresh_from_db()
        self.assertEqual(self.enrollment.course, self.course_a)

    def test_transfer_keeps_payments_and_recalculates_installments(self):
        self.course_b.enrollment_fee = 50
        self.course_b.save()
        self.assertEqual(self.transfer(self.course_b, use_fees=True).status_code, 302)
        self.enrollment.refresh_from_db()
        self.fee.refresh_from_db()
        self.monthly.refresh_from_db()
        self.payment.refresh_from_db()
        self.assertEqual(self.enrollment.course, self.course_b)
        self.assertEqual((self.enrollment.enrollment_fee, self.enrollment.monthly_fee), (50, 70))
        self.assertEqual((self.fee.amount, self.monthly.amount), (50, 70))
        self.assertEqual(self.fee.due_date, date(2026, 11, 2))
        self.assertEqual(self.monthly.concept, "Mensualidad inicial (2026-11)")
        self.assertEqual(self.payment.status, Payment.Status.CONFIRMED)
        self.assertEqual(self.fee.balance, 10)

    def test_transfer_can_keep_current_fees(self):
        self.assertEqual(self.transfer(self.course_b, use_fees=False).status_code, 302)
        self.enrollment.refresh_from_db()
        self.assertEqual(self.enrollment.course, self.course_b)
        self.assertEqual((self.enrollment.enrollment_fee, self.enrollment.monthly_fee), (40, 60))

    def test_transfer_requires_permission_and_hides_taken_courses(self):
        Enrollment.objects.create(campus=self.campus, student=self.student, course=self.course_b)
        form = self.client.get(reverse("enrollment-transfer", args=[self.enrollment.pk])).context[
            "form"
        ]
        self.assertNotIn(self.course_b, form.fields["course"].queryset)
        self.assertNotIn(self.course_a, form.fields["course"].queryset)

        clerk = create_user(username="clerk", password="x-Test-pass-123")
        clerk.campuses.add(self.campus)
        self.client.force_login(clerk)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        response = self.client.get(reverse("enrollment-transfer", args=[self.enrollment.pk]))
        self.assertEqual(response.status_code, 403)

    def test_student_trash_restore_brings_back_enrollments_installments_and_receipts(self):
        self.client.post(reverse("student-delete", args=[self.student.pk]))
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.status, Payment.Status.VOID)
        self.assertContains(self.client.get(reverse("trash")), "ANA PÉREZ")

        self.client.post(reverse("student-restore", args=[self.student.pk]))
        self.student.refresh_from_db()
        self.enrollment.refresh_from_db()
        self.payment.refresh_from_db()
        self.assertTrue(self.student.active)
        self.assertTrue(self.enrollment.active)
        self.assertEqual(self.payment.status, Payment.Status.CONFIRMED)
        self.assertEqual(
            Installment.objects.filter(enrollment=self.enrollment, active=True).count(), 2
        )
        self.assertEqual(list(self.client.get(reverse("trash")).context["students"]), [])

    def test_enrollment_restore_keeps_manually_voided_receipts_void(self):
        manual = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="R-2",
            amount=10,
            status=Payment.Status.VOID,
        )
        self.client.post(reverse("enrollment-delete", args=[self.enrollment.pk]))
        self.client.post(reverse("enrollment-restore", args=[self.enrollment.pk]))
        self.enrollment.refresh_from_db()
        self.payment.refresh_from_db()
        manual.refresh_from_db()
        self.assertTrue(self.enrollment.active)
        self.assertEqual(self.payment.status, Payment.Status.CONFIRMED)
        self.assertEqual(manual.status, Payment.Status.VOID)

    def test_enrollment_restore_requires_active_student(self):
        self.client.post(reverse("student-delete", args=[self.student.pk]))
        self.client.post(reverse("enrollment-restore", args=[self.enrollment.pk]))
        self.enrollment.refresh_from_db()
        self.assertFalse(self.enrollment.active)


class StudentEnrollmentCrudTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede CRUD", active=True)
        self.user = create_user(
            username="manager",
            password="secret123",
            manage_students=True,
            register_payments=True,
        )
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.student = Student.objects.create(
            campus=self.campus,
            full_name="Estudiante de prueba",
            document_number="11112222",
            phone="987654321",
        )
        self.student_phone = StudentPhone.objects.create(student=self.student, number="987654321")
        self.course = Course.objects.create(
            campus=self.campus,
            name="Curso de prueba",
            start_date=date(2026, 10, 14),
            enrollment_fee=40,
            monthly_fee=60,
        )
        self.enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=self.student,
            course=self.course,
            enrollment_fee=40,
            monthly_fee=60,
        )

    def add_payment_history(self):
        installment = Installment.objects.create(
            enrollment=self.enrollment,
            concept="Cuota de prueba",
            amount=100,
            due_date=date(2026, 10, 14),
        )
        payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="CRUD-001",
            amount=40,
        )
        allocation = PaymentAllocation.objects.create(
            payment=payment,
            installment=installment,
            amount=40,
        )
        return installment, payment, allocation

    def phone_formset_data(self, numbers):
        existing = list(self.student.phones.order_by("pk"))
        form_count = max(len(existing), len(numbers), 1)
        data = {
            "phones-TOTAL_FORMS": str(form_count),
            "phones-INITIAL_FORMS": str(len(existing)),
            "phones-MIN_NUM_FORMS": "1",
            "phones-MAX_NUM_FORMS": "1000",
        }
        for index, phone in enumerate(existing):
            data[f"phones-{index}-id"] = str(phone.pk)
        for index, number in enumerate(numbers):
            data[f"phones-{index}-number"] = number
        return data

    def test_student_and_enrollment_can_be_updated(self):
        response = self.client.get(reverse("student-update", args=[self.student.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(
            reverse("student-update", args=[self.student.pk]),
            {
                "last_name": "Apellidos actualizados",
                "first_name": "Nombres",
                "document_number": self.student.document_number,
                "active": "on",
                **self.phone_formset_data(["987654321"]),
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.student.refresh_from_db()
        self.assertEqual(self.student.full_name, "APELLIDOS ACTUALIZADOS NOMBRES")

        response = self.client.get(reverse("enrollment-update", args=[self.enrollment.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(
            reverse("enrollment-update", args=[self.enrollment.pk]),
            {
                "course": self.course.pk,
                "student": self.student.pk,
                "record_code": "202601",
                "promoter": [],
                "enrollment_date": self.enrollment.enrollment_date.isoformat(),
                "enrollment_fee": "40.00",
                "monthly_fee": "60.00",
                "active": "on",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.enrollment.refresh_from_db()
        self.assertEqual(self.enrollment.record_code, 202601)

    def test_student_and_enrollment_views_reject_other_campus_records(self):
        other_campus = Campus.objects.create(name="Sede ajena")
        other_student = Student.objects.create(campus=other_campus, full_name="Alumno ajeno")
        other_course = Course.objects.create(
            campus=other_campus, name="Curso ajeno", start_date=date(2026, 10, 14)
        )
        other_enrollment = Enrollment.objects.create(
            campus=other_campus, student=other_student, course=other_course
        )

        requests = (
            ("get", reverse("student-update", args=[other_student.pk])),
            ("post", reverse("student-delete", args=[other_student.pk])),
            ("get", reverse("enrollment-update", args=[other_enrollment.pk])),
            ("post", reverse("enrollment-delete", args=[other_enrollment.pk])),
            ("get", reverse("enrollment-transfer", args=[other_enrollment.pk])),
            ("post", reverse("student-restore", args=[other_student.pk])),
            ("post", reverse("enrollment-restore", args=[other_enrollment.pk])),
        )
        for method, url in requests:
            with self.subTest(url=url, method=method):
                self.assertEqual(getattr(self.client, method)(url).status_code, 404)

    def test_student_views_require_the_assigned_permission_group(self):
        revoke(self.user, "manage_students")

        self.assertEqual(self.client.get(reverse("student-list")).status_code, 403)
        self.assertEqual(self.client.get(reverse("student-create")).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("student-update", args=[self.student.pk])).status_code, 403
        )

    def test_enrollment_record_code_accepts_only_numbers_or_empty(self):
        response = self.client.post(
            reverse("enrollment-update", args=[self.enrollment.pk]),
            {
                "course": self.course.pk,
                "student": self.student.pk,
                "record_code": "F-2026-01",
                "promoter": [],
                "enrollment_date": self.enrollment.enrollment_date.isoformat(),
                "enrollment_fee": "40.00",
                "monthly_fee": "60.00",
                "active": "on",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ingrese solo números en el código de ficha.")
        self.enrollment.refresh_from_db()
        self.assertEqual(self.enrollment.record_code, "")

    def test_ficha_exists_only_when_record_code_is_given(self):
        base = {
            "course": self.course.pk,
            "student": self.student.pk,
            "promoter": [],
            "enrollment_date": self.enrollment.enrollment_date.isoformat(),
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        url = reverse("enrollment-update", args=[self.enrollment.pk])
        for data in ({"record_code": "0"}, {"record_code": "-3"}, {"record_code": "abc"}):
            response = self.client.post(url, {**base, **data})
            self.assertEqual(response.status_code, 200, data)
            self.assertTrue(response.context["form"].errors.get("record_code"), data)

        # Sin código no se crea ficha.
        self.assertEqual(self.client.post(url, {**base, "record_code": ""}).status_code, 302)
        self.assertFalse(Ficha.objects.filter(enrollment=self.enrollment).exists())
        response = self.client.post(url, {**base, "record_code": "15"})
        self.assertEqual(response.status_code, 302)
        self.enrollment.refresh_from_db()
        self.assertTrue(self.enrollment.has_record)
        self.assertEqual(self.enrollment.record_code, 15)

    def test_ficha_image_is_optional_and_can_be_replaced_and_removed(self):
        import io
        import shutil
        import tempfile

        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from PIL import Image

        def png(name):
            buffer = io.BytesIO()
            Image.new("RGB", (4, 4), "red").save(buffer, "PNG")
            return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/png")

        base = {
            "course": self.course.pk,
            "student": self.student.pk,
            "record_code": "77",
            "promoter": [],
            "enrollment_date": self.enrollment.enrollment_date.isoformat(),
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        url = reverse("enrollment-update", args=[self.enrollment.pk])
        media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media, ignore_errors=True)
        with override_settings(MEDIA_ROOT=media):
            self.assertEqual(self.client.post(url, base).status_code, 302)
            ficha = Ficha.objects.get(enrollment=self.enrollment)
            self.assertFalse(ficha.image)

            with self.captureOnCommitCallbacks(execute=True):
                response = self.client.post(url, {**base, "record_image": png("a.png")})
            self.assertEqual(response.status_code, 302)
            ficha.refresh_from_db()
            first = ficha.image.name
            self.assertTrue(first.startswith("fichas/"))
            self.assertTrue(first.endswith(".jpg"))
            with ficha.image.open("rb") as handle:
                self.assertEqual(Image.open(handle).format, "JPEG")

            big = io.BytesIO()
            import random

            noise = Image.frombytes(
                "RGB", (400, 400), bytes(random.randrange(256) for _ in range(400 * 400 * 3))
            )
            noise.save(big, "PNG")
            sizes = {}
            for quality in ("90", "60"):
                upload = SimpleUploadedFile("n.png", big.getvalue(), content_type="image/png")
                with self.captureOnCommitCallbacks(execute=True):
                    self.client.post(
                        url, {**base, "record_image": upload, "record_image_quality": quality}
                    )
                ficha.refresh_from_db()
                sizes[quality] = ficha.image.size
            self.assertLess(sizes["60"], sizes["90"])
            first = ficha.image.name

            with self.captureOnCommitCallbacks(execute=True):
                self.client.post(url, {**base, "record_image": png("b.png")})
            ficha.refresh_from_db()
            self.assertNotEqual(ficha.image.name, first)
            self.assertFalse(ficha.image.storage.exists(first))

            self.client.post(url, base)
            ficha.refresh_from_db()
            self.assertTrue(ficha.image)

            with self.captureOnCommitCallbacks(execute=True):
                self.client.post(url, {**base, "record_image-clear": "on"})
            ficha.refresh_from_db()
            self.assertFalse(ficha.image)

    def test_clearing_record_code_deletes_ficha_and_image_requires_code(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        Ficha.objects.create(
            enrollment=self.enrollment,
            campus=self.campus,
            date=self.enrollment.enrollment_date,
            code="8",
        )
        base = {
            "course": self.course.pk,
            "student": self.student.pk,
            "promoter": [],
            "enrollment_date": self.enrollment.enrollment_date.isoformat(),
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        url = reverse("enrollment-update", args=[self.enrollment.pk])
        response = self.client.post(
            url, {**base, "record_image": SimpleUploadedFile("x.png", b"x", "image/png")}
        )
        self.assertIn("record_image", response.context["form"].errors)
        self.assertEqual(self.client.post(url, {**base, "record_code": ""}).status_code, 302)
        self.assertFalse(Ficha.objects.filter(enrollment=self.enrollment).exists())

    def test_changing_enrollment_date_updates_ficha_date(self):
        ficha = Ficha.objects.create(
            enrollment=self.enrollment,
            campus=self.campus,
            date=self.enrollment.enrollment_date,
            code="9",
        )
        self.enrollment.enrollment_date = date(2026, 1, 5)
        self.enrollment.save()
        ficha.refresh_from_db()
        self.assertEqual(ficha.date, date(2026, 1, 5))

    def test_editing_enrollment_recalculates_installments(self):
        data = {
            "course": self.course.pk,
            "student": self.student.pk,
            "promoter": [],
            "enrollment_date": self.enrollment.enrollment_date.isoformat(),
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        url = reverse("enrollment-update", args=[self.enrollment.pk])
        due = self.enrollment.enrollment_date
        matricula = Installment.objects.create(
            enrollment=self.enrollment, concept="Matrícula", amount=30, due_date=due
        )
        monthly = Installment.objects.create(
            enrollment=self.enrollment,
            concept="Mensualidad inicial (2026-01)",
            amount=50,
            due_date=due,
        )

        response = self.client.post(url, data)

        self.assertEqual(response.status_code, 302)
        matricula.refresh_from_db()
        monthly.refresh_from_db()
        self.assertEqual(matricula.amount, 40)
        self.assertEqual(monthly.amount, 60)

    def test_fee_cannot_drop_below_amount_already_paid(self):
        data = {
            "course": self.course.pk,
            "student": self.student.pk,
            "promoter": [],
            "enrollment_date": self.enrollment.enrollment_date.isoformat(),
            "enrollment_fee": "10.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        due = self.enrollment.enrollment_date
        matricula = Installment.objects.create(
            enrollment=self.enrollment, concept="Matrícula", amount=40, due_date=due
        )
        payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=User.objects.first(),
            receipt_number="900",
            amount=25,
        )
        PaymentAllocation.objects.create(payment=payment, installment=matricula, amount=25)

        response = self.client.post(reverse("enrollment-update", args=[self.enrollment.pk]), data)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["form"].errors.get("enrollment_fee"))
        matricula.refresh_from_db()
        self.assertEqual(matricula.amount, 40)

    def test_student_form_inputs_convert_names_to_uppercase(self):
        response = self.client.get(reverse("student-create"))

        self.assertContains(response, 'name="last_name"')
        self.assertContains(response, "<input", count=None)
        self.assertContains(response, 'data-uppercase style="text-transform:uppercase"', count=2)

    def test_student_creation_requires_separate_last_and_first_names(self):
        response = self.client.post(
            reverse("student-create"),
            {
                "last_name": "Pérez",
                "first_name": "Ana",
                "document_number": "12345678",
                "active": "on",
                "phones-TOTAL_FORMS": "1",
                "phones-INITIAL_FORMS": "0",
                "phones-MIN_NUM_FORMS": "1",
                "phones-MAX_NUM_FORMS": "1000",
                "phones-0-number": "987654321",
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(full_name="PÉREZ ANA")
        self.assertEqual(student.last_name, "PÉREZ")
        self.assertEqual(student.first_name, "ANA")
        self.assertEqual(student.phone, "987654321")
        self.assertEqual(list(student.phones.values_list("number", flat=True)), ["987654321"])

        invalid_response = self.client.post(
            reverse("student-create"),
            {
                "last_name": "",
                "first_name": "Ana",
                "active": "on",
                "phones-TOTAL_FORMS": "1",
                "phones-INITIAL_FORMS": "0",
                "phones-MIN_NUM_FORMS": "1",
                "phones-MAX_NUM_FORMS": "1000",
                "phones-0-number": "987654321",
            },
        )
        self.assertEqual(invalid_response.status_code, 200)
        self.assertContains(invalid_response, "Este campo es obligatorio.")

    def test_student_names_are_unique_per_campus(self):
        other_campus = Campus.objects.create(name="Otra sede", active=True)
        same_name_at_other_campus = Student.objects.create(
            campus=other_campus,
            full_name=self.student.full_name,
            phone="987654321",
        )
        self.assertEqual(same_name_at_other_campus.full_name, self.student.full_name)

        duplicate_at_same_campus = Student(
            campus=self.campus,
            full_name=self.student.full_name,
            phone="987654321",
        )
        with self.assertRaises(ValidationError):
            duplicate_at_same_campus.full_clean()

    def test_pending_course_enrollments_allow_repeated_records_without_a_fake_course(self):
        for _ in range(2):
            Enrollment(
                campus=self.campus,
                student=self.student,
                course=None,
                enrollment_fee=10,
                monthly_fee=10,
            ).full_clean()
            Enrollment.objects.create(campus=self.campus, student=self.student, course=None)
        self.assertEqual(
            Enrollment.objects.filter(student=self.student, course__isnull=True).count(), 2
        )
        response = self.client.get(reverse("enrollment-create"), {"student": self.student.pk})
        self.assertContains(response, "PENDIENTE DE ASIGNACIÓN")
        self.assertNotContains(response, "SIN CURSO")
        self.assertNotContains(response, f'<option value="{self.course.pk}"')
        with self.assertRaises(ValidationError):
            Enrollment(campus=self.campus, student=self.student, course=self.course).full_clean(
                exclude=["enrollment_fee", "monthly_fee"]
            )

    def test_default_promoters_cannot_be_combined_with_others(self):
        other = Promoter.objects.create(name="Otro Promotor")
        other.campuses.add(self.campus)
        particular, _ = Promoter.objects.get_or_create(name="PARTICULAR")
        particular.campuses.add(self.campus)
        base = {
            "course": self.course.pk,
            "student": self.student.pk,
            "enrollment_date": self.enrollment.enrollment_date.isoformat(),
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        url = reverse("enrollment-update", args=[self.enrollment.pk])
        response = self.client.get(url)
        labels = [c[1] for c in response.context["form"].fields["promoter"].choices]
        self.assertEqual(labels[-2:], ["NN", "PARTICULAR"])
        self.assertEqual(labels[0], "OTRO PROMOTOR")
        response = self.client.post(url, {**base, "promoter": [particular.pk, other.pk]})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no se pueden combinar con otros promotores")
        response = self.client.post(url, {**base, "promoter": [particular.pk]})
        self.assertEqual(response.status_code, 302)
        response = self.client.post(url, {**base, "promoter": [other.pk]})
        self.assertEqual(response.status_code, 302)

    def test_enrollment_create_lists_student_current_enrollments(self):
        response = self.client.get(reverse("enrollment-create"), {"student": self.student.pk})
        self.assertContains(response, "Cursos en los que ya está matriculado")
        self.assertContains(response, self.course.name)

        other = Student.objects.create(campus=self.campus, full_name="Sin matriculas")
        response = self.client.get(reverse("enrollment-create"), {"student": other.pk})
        self.assertNotContains(response, "Cursos en los que ya está matriculado")

    def test_record_code_cannot_repeat_on_same_enrollment_date(self):
        self.enrollment.enrollment_date = date(2026, 9, 29)
        self.enrollment.save()
        Ficha.objects.create(
            enrollment=self.enrollment,
            campus=self.campus,
            date=self.enrollment.enrollment_date,
            code="500",
        )
        other = Student.objects.create(campus=self.campus, full_name="Alumno ficha repetida")
        other_course = Course.objects.create(
            campus=self.campus, name="Otro curso", enrollment_fee=10, monthly_fee=10
        )
        data = {
            "course": other_course.pk,
            "student": other.pk,
            "record_code": "500",
            "promoter": [],
            "enrollment_date": "2026-09-29",
            "enrollment_fee": "10.00",
            "monthly_fee": "10.00",
            "active": "on",
        }
        url = reverse("enrollment-create") + f"?student={other.pk}"
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 200)
        self.assertIn("record_code", response.context["form"].errors)
        self.assertFalse(Enrollment.objects.filter(student=other).exists())
        data["enrollment_date"] = "2026-09-30"
        self.assertEqual(self.client.post(url, data).status_code, 302)
        self.assertTrue(Enrollment.objects.filter(student=other).exists())

    def test_enrollment_form_offers_pending_course_as_an_empty_choice(self):
        available_student = Student.objects.create(
            campus=self.campus, full_name="Estudiante pendiente"
        )
        response = self.client.get(reverse("enrollment-create"), {"student": available_student.pk})
        form = response.context["form"]
        self.assertIsNone(form.initial.get("course"))
        self.assertEqual(form.fields["course"].queryset.first(), self.course)
        self.assertEqual(str(list(form.fields["course"].choices)[0][0]), "")
        self.assertEqual(form.fields["course"].empty_label, "PENDIENTE DE ASIGNACIÓN")
        self.assertNotContains(response, "SIN CURSO")

    def test_enrollment_without_course_forces_zero_fees(self):
        new_student = Student.objects.create(campus=self.campus, full_name="Alumno por asignar")
        response = self.client.post(
            reverse("enrollment-create") + f"?student={new_student.pk}",
            {
                "course": "",
                "student": new_student.pk,
                "promoter": [],
                "enrollment_date": "2026-09-29",
                "enrollment_fee": "99.00",
                "monthly_fee": "",
                "active": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        enrollment = Enrollment.objects.get(student=new_student, course__isnull=True)
        self.assertEqual(enrollment.enrollment_fee, 0)
        self.assertEqual(enrollment.monthly_fee, 0)
        self.assertFalse(Installment.objects.filter(enrollment=enrollment).exists())

    def test_enrollment_without_course_creates_pending_enrollment_in_one_click(self):
        new_student = Student.objects.create(campus=self.campus, full_name="Alumno sin curso")
        self.assertContains(self.client.get(reverse("enrollment-create")), "Matrícula sin curso")
        response = self.client.post(
            reverse("enrollment-create") + f"?student={new_student.pk}",
            {"student": new_student.pk, "pending_course": "1"},
        )
        self.assertEqual(response.status_code, 302)
        enrollment = Enrollment.objects.get(student=new_student)
        self.assertIsNone(enrollment.course)
        self.assertEqual(enrollment.course_label, "PENDIENTE DE ASIGNACIÓN")
        self.assertEqual(enrollment.enrollment_fee, 0)
        self.assertFalse(Installment.objects.filter(enrollment=enrollment).exists())

    def test_pending_course_enrollment_is_not_payable(self):
        pending = Enrollment.objects.create(campus=self.campus, student=self.student, course=None)
        Installment.objects.create(
            enrollment=pending,
            concept="Cuota incorrecta",
            amount=10,
            due_date=date(2026, 10, 14),
        )

        from payments.forms import pending_enrollments_for_campus

        self.assertNotIn(pending, pending_enrollments_for_campus(self.campus))

    def test_enrollment_without_student_reserves_course_seat(self):
        url = reverse("enrollment-create") + "?sin_estudiante=1"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        form = response.context["form"]
        self.assertNotIn("student", form.fields)
        self.assertTrue(form.fields["course"].required)
        self.assertContains(self.client.get(reverse("enrollment-create")), "sin_estudiante=1")

        data = {
            "sin_estudiante": "1",
            "course": self.course.pk,
            "promoter": [],
            "enrollment_date": "2026-09-29",
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        bad = self.client.post(url, {**data, "course": ""})
        self.assertEqual(bad.status_code, 200)
        self.assertIn("course", bad.context["form"].errors)

        self.assertEqual(self.client.post(url, data).status_code, 302)
        seat = Enrollment.objects.get(course=self.course, student__isnull=True)
        self.assertEqual(seat.student_label, "SIN ESTUDIANTE")
        self.assertEqual(Installment.objects.filter(enrollment=seat).count(), 2)
        self.assertContains(self.client.get(reverse("enrollment-list")), "SIN ESTUDIANTE")

        from payments.forms import pending_enrollments_for_campus

        self.assertNotIn(seat, pending_enrollments_for_campus(self.campus))

    def test_enrollment_shows_remaining_seats_message(self):
        self.course.capacity = 5
        self.course.save()
        url = reverse("enrollment-create") + "?sin_estudiante=1"
        data = {
            "sin_estudiante": "1",
            "course": self.course.pk,
            "promoter": [],
            "enrollment_date": "2026-09-29",
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        response = self.client.post(url, data, follow=True)
        texts = [str(m) for m in response.context["messages"]]
        self.assertTrue(any("Quedan 3 vacantes disponibles" in t for t in texts), texts)

    def test_enrollment_without_student_can_be_assigned_later(self):
        seat = Enrollment.objects.create(
            campus=self.campus, course=self.course, enrollment_fee=40, monthly_fee=60
        )
        new_student = Student.objects.create(campus=self.campus, full_name="Alumno asignado")
        url = reverse("enrollment-update", args=[seat.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertContains(self.client.get(reverse("enrollment-list")), "Pendiente")
        self.assertContains(self.client.get(reverse("enrollment-list")), "Asignar Estudiante")
        data = {
            "course": self.course.pk,
            "student": new_student.pk,
            "promoter": [],
            "enrollment_date": "2026-09-29",
            "enrollment_fee": "40.00",
            "monthly_fee": "60.00",
            "active": "on",
        }
        from django.contrib.auth.models import Group, Permission

        group = Group.objects.create(name="solo_matricula")
        group.permissions.add(
            *Permission.objects.filter(
                content_type__app_label="students",
                codename__in=["view_enrollment", "add_enrollment", "view_student", "view_course"],
            )
        )
        clerk = create_user(username="clerk", password="Segura12345!")
        clerk.groups.add(group)
        clerk.campuses.add(self.campus)
        self.client.force_login(clerk)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.assertFalse(clerk.has_perm("students.change_enrollment"))
        self.assertFalse(clerk.has_perm("core.change_enrollment_course"))

        self.assertEqual(self.client.post(url, data).status_code, 302)
        seat.refresh_from_db()
        self.assertEqual(seat.student, new_student)

        # Una matrícula completa sigue exigiendo permiso de edición.
        self.assertEqual(
            self.client.get(reverse("enrollment-update", args=[self.enrollment.pk])).status_code,
            403,
        )

        pending = Enrollment.objects.create(campus=self.campus, student=self.student, course=None)
        free = Course.objects.create(
            campus=self.campus,
            name="Libre",
            start_date=date(2026, 10, 14),
            enrollment_fee=10,
            monthly_fee=10,
        )
        self.assertEqual(
            self.client.get(reverse("enrollment-update", args=[pending.pk])).status_code, 200
        )
        response = self.client.post(
            reverse("enrollment-update", args=[pending.pk]),
            {
                "course": free.pk,
                "student": self.student.pk,
                "promoter": [],
                "enrollment_date": "2026-09-29",
                "enrollment_fee": "10.00",
                "monthly_fee": "10.00",
                "active": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        pending.refresh_from_db()
        self.assertEqual(pending.course, free)

    def test_pending_enrollment_locks_student_and_hides_taken_courses(self):
        free = Course.objects.create(
            campus=self.campus,
            name="Libre",
            start_date=date(2026, 10, 14),
            enrollment_fee=10,
            monthly_fee=10,
        )
        pending = Enrollment.objects.create(campus=self.campus, student=self.student, course=None)
        response = self.client.get(reverse("enrollment-update", args=[pending.pk]))
        form = response.context["form"]
        self.assertTrue(form.fields["student"].disabled)
        courses = form.fields["course"].queryset
        self.assertIn(free, courses)
        self.assertNotIn(self.course, courses)
        self.assertContains(self.client.get(reverse("enrollment-list")), "Asignar curso")

    def test_available_students_ignore_seats_without_student(self):
        seat = Enrollment.objects.create(
            campus=self.campus, course=self.course, enrollment_fee=40, monthly_fee=60
        )
        other = Student.objects.create(campus=self.campus, full_name="Alumno libre")
        response = self.client.get(
            reverse("course-students-available"),
            {"course_id": self.course.pk, "enrollment_id": seat.pk},
        )
        self.assertIn(other.pk, [s["id"] for s in response.json()["students"]])
        edit = self.client.get(reverse("enrollment-update", args=[seat.pk]))
        self.assertIn(other, edit.context["form"].fields["student"].queryset)

    def test_enrollment_with_student_cannot_lose_it(self):
        response = self.client.post(
            reverse("enrollment-update", args=[self.enrollment.pk]),
            {
                "course": self.course.pk,
                "student": "",
                "promoter": [],
                "enrollment_date": "2026-09-29",
                "enrollment_fee": "40.00",
                "monthly_fee": "60.00",
                "active": "on",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("student", response.context["form"].errors)

    def test_enrollment_creation_saves_user_selected_date(self):
        new_student = Student.objects.create(
            campus=self.campus,
            full_name="Alumno con fecha elegida",
        )
        response = self.client.get(reverse("enrollment-create"))
        self.assertContains(response, "selecciona un estudiante")
        self.assertContains(response, new_student.full_name)
        self.assertContains(response, 'id="student-search"')
        filtered = self.client.get(reverse("enrollment-create"), {"q": "fecha elegida"})
        self.assertContains(filtered, new_student.full_name)
        self.assertNotContains(filtered, self.student.full_name)
        response = self.client.get(reverse("enrollment-create"), {"student": new_student.pk})
        self.assertContains(response, 'type="date"')
        self.assertContains(response, f'value="{timezone.localdate():%Y-%m-%d}"')
        response = self.client.post(
            reverse("enrollment-create"),
            {
                "course": self.course.pk,
                "student": new_student.pk,
                "record_code": "202602",
                "promoter": [],
                "enrollment_date": "2026-09-29",
                "enrollment_fee": "25.00",
                "monthly_fee": "35.00",
                "active": "on",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        created_enrollment = Enrollment.objects.get(student=new_student, course=self.course)
        self.assertEqual(created_enrollment.enrollment_date, date(2026, 9, 29))
        self.assertContains(response, "29/09/2026")

    def test_search_ignores_accents_and_case(self):
        Student.objects.create(campus=self.campus, full_name="Canción Ñandú")
        for term in ("cancion", "CANCIÓN", "CANcion", "nandu"):
            response = self.client.get(reverse("student-list"), {"q": term})
            self.assertContains(response, "CANCIÓN ÑANDÚ", msg_prefix=term)

    def test_list_pages_use_live_filtering(self):
        for name in (
            "student-list",
            "enrollment-list",
            "course-student-tables",
            "payment-list",
            "pending-debts",
            "payment-create",
        ):
            response = self.client.get(reverse(name))
            self.assertContains(response, "data-live-filter", msg_prefix=name)
            self.assertContains(response, "data-live-region", msg_prefix=name)

    def test_enrollment_edit_preserves_custom_fees_instead_of_course_fees(self):
        self.enrollment.enrollment_fee = 25
        self.enrollment.monthly_fee = 55
        self.enrollment.save(update_fields=["enrollment_fee", "monthly_fee"])

        response = self.client.get(reverse("enrollment-update", args=[self.enrollment.pk]))

        self.assertEqual(response.status_code, 200)
        form = response.context["form"]
        self.assertEqual(str(form["enrollment_fee"].value()), "25.00")
        self.assertEqual(str(form["monthly_fee"].value()), "55.00")
        self.assertNotContains(response, "updateCourseFees();")

    def test_enrollment_can_have_multiple_promoters(self):
        first_promoter = Promoter.objects.create(name="Primer promotor", active=True)
        second_promoter = Promoter.objects.create(name="Segundo promotor", active=True)
        response = self.client.get(reverse("enrollment-update", args=[self.enrollment.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'type="checkbox"')
        self.assertContains(response, "PRIMER PROMOTOR")
        self.assertContains(response, "SEGUNDO PROMOTOR")
        response = self.client.post(
            reverse("enrollment-update", args=[self.enrollment.pk]),
            {
                "course": self.course.pk,
                "student": self.student.pk,
                "record_code": "202603",
                "promoter": [first_promoter.pk, second_promoter.pk],
                "enrollment_date": self.enrollment.enrollment_date.isoformat(),
                "enrollment_fee": "40.00",
                "monthly_fee": "60.00",
                "active": "on",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertSetEqual(
            set(self.enrollment.promoter.values_list("pk", flat=True)),
            {first_promoter.pk, second_promoter.pk},
        )

    def test_deleting_student_deactivates_records_and_voids_receipt(self):
        installment, payment, allocation = self.add_payment_history()
        response = self.client.get(reverse("student-delete", args=[self.student.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(reverse("student-delete", args=[self.student.pk]), follow=True)
        self.assertEqual(response.status_code, 200)
        self.student.refresh_from_db()
        self.enrollment.refresh_from_db()
        installment.refresh_from_db()
        payment.refresh_from_db()
        self.assertFalse(self.student.active)
        self.assertFalse(self.enrollment.active)
        self.assertFalse(installment.active)
        self.assertEqual(payment.status, Payment.Status.VOID)
        self.assertTrue(Payment.objects.filter(pk=payment.pk).exists())
        self.assertTrue(PaymentAllocation.objects.filter(pk=allocation.pk).exists())

    def test_deleting_enrollment_deactivates_and_preserves_history(self):
        installment, payment, allocation = self.add_payment_history()
        response = self.client.get(reverse("enrollment-delete", args=[self.enrollment.pk]))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(
            reverse("enrollment-delete", args=[self.enrollment.pk]), follow=True
        )
        self.assertEqual(response.status_code, 200)
        self.enrollment.refresh_from_db()
        installment.refresh_from_db()
        payment.refresh_from_db()
        self.assertFalse(self.enrollment.active)
        self.assertFalse(installment.active)
        self.assertEqual(payment.status, Payment.Status.VOID)
        self.assertTrue(PaymentAllocation.objects.filter(pk=allocation.pk).exists())

    def test_student_and_enrollment_lists_link_to_payment_form(self):
        Installment.objects.create(
            enrollment=self.enrollment,
            concept="Saldo pendiente",
            amount=100,
            due_date=date(2026, 10, 14),
        )
        response = self.client.get(reverse("student-list"))
        self.assertContains(
            response,
            f'href="{reverse("payment-create")}?student={self.student.pk}"',
        )
        self.assertContains(
            response,
            f'href="{reverse("enrollment-create")}?student={self.student.pk}"',
        )
        self.assertContains(response, 'aria-label="Editar estudiante"')
        self.assertContains(response, 'aria-label="Dar de baja estudiante"')
        response = self.client.get(reverse("enrollment-list"))
        self.assertContains(
            response,
            f'href="{reverse("payment-create-for", args=[self.enrollment.pk])}"',
        )
        self.assertContains(response, 'aria-label="Pagar matrícula"')
        self.assertContains(response, 'aria-label="Editar matrícula"')
        self.assertContains(response, 'aria-label="Dar de baja matrícula"')
        response = self.client.get(reverse("payment-create"), {"student": self.student.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([e.pk for e in response.context["enrollments"]], [self.enrollment.pk])

    def test_student_list_shows_active_enrollment_count_and_highlights_zero(self):
        student_without_enrollments = Student.objects.create(
            campus=self.campus,
            full_name="Estudiante sin matrícula",
        )
        Enrollment.objects.create(
            campus=self.campus,
            student=student_without_enrollments,
            course=self.course,
            active=False,
        )

        response = self.client.get(reverse("student-list"))
        students = {student.pk: student for student in response.context["students"]}

        self.assertEqual(students[self.student.pk].enrollment_count, 1)
        self.assertEqual(students[student_without_enrollments.pk].enrollment_count, 0)
        self.assertEqual(students[self.student.pk].row_number, 1)
        self.assertContains(response, "Cursos matriculados")
        self.assertContains(response, ">N.º ↓</a>")
        self.assertContains(response, "chip red darken-1")
        self.assertNotContains(response, "<th>Sede</th>")

        by_name = self.client.get(reverse("student-list"), {"sort": "name", "direction": "desc"})
        self.assertEqual(by_name.context["students"][0].pk, student_without_enrollments.pk)
        by_enrollments = self.client.get(
            reverse("student-list"), {"sort": "enrollments", "direction": "desc"}
        )
        self.assertEqual(by_enrollments.context["students"][0].pk, self.student.pk)
        by_row_number = self.client.get(
            reverse("student-list"), {"sort": "row_number", "direction": "desc"}
        )
        self.assertEqual(by_row_number.context["students"][0].pk, student_without_enrollments.pk)
        self.assertContains(response, "?sort=row_number&amp;direction=asc")

    def test_student_name_links_to_only_that_students_enrollments(self):
        other_student = Student.objects.create(
            campus=self.campus,
            full_name="Otra alumna",
        )
        other_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=other_student,
            course=self.course,
        )

        response = self.client.get(reverse("enrollment-list"))
        self.assertContains(
            response,
            f'href="{reverse("student-enrollment-list", args=[self.student.pk])}">{self.student.full_name}</a>',
        )

        response = self.client.get(reverse("student-enrollment-list", args=[self.student.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'name="student" value="{self.student.full_name}"')
        self.assertEqual(
            {enrollment.pk for enrollment in response.context["enrollments"]},
            {self.enrollment.pk},
        )
        self.assertNotIn(
            other_enrollment.pk, {enrollment.pk for enrollment in response.context["enrollments"]}
        )

    def test_course_promoter_and_enrollment_date_links_filter_exactly(self):
        promoter = Promoter.objects.create(name="Promotor seleccionado", active=True)
        self.enrollment.enrollment_date = date(2026, 9, 29)
        self.enrollment.save(update_fields=["enrollment_date"])
        self.enrollment.promoter.add(promoter)
        other_student = Student.objects.create(campus=self.campus, full_name="Otra estudiante")
        other_course = Course.objects.create(
            campus=self.campus,
            name=self.course.name,
            schedule="Miércoles",
            group="B",
        )
        other_promoter = Promoter.objects.create(name="Otro promotor", active=True)
        other_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=other_student,
            course=other_course,
            enrollment_date=date(2026, 10, 1),
        )
        other_enrollment.promoter.add(other_promoter)

        response = self.client.get(reverse("enrollment-list"))
        self.assertContains(response, f"?course_id={self.course.pk}")
        self.assertContains(response, f"?promoter_id={promoter.pk}")
        self.assertContains(response, "?enrollment_date=2026-09-29")

        filters = [
            ({"course_id": self.course.pk}, "course", self.course.name),
            ({"promoter_id": promoter.pk}, "promoter", promoter.name),
            ({"enrollment_date": "2026-09-29"}, "enrollment_date_from", "2026-09-29"),
        ]
        for query, input_name, expected_value in filters:
            with self.subTest(query=query):
                response = self.client.get(reverse("enrollment-list"), query)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    {enrollment.pk for enrollment in response.context["enrollments"]},
                    {self.enrollment.pk},
                )
                self.assertContains(response, f'name="{input_name}" value="{expected_value}"')
                if "enrollment_date" in query:
                    self.assertContains(response, 'name="enrollment_date_to" value="2026-09-29"')

    def test_enrollment_list_shows_costs_and_outstanding_debt(self):
        self.add_payment_history()
        Installment.objects.create(
            enrollment=self.enrollment,
            concept="Cuota adicional",
            amount=30,
            due_date=date(2026, 11, 14),
        )

        response = self.client.get(reverse("enrollment-list"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Costo de inscripción")
        self.assertContains(response, "Mensualidad")
        self.assertContains(response, "Pago total")
        self.assertContains(response, "Pagado")
        self.assertContains(response, "Deuda")
        self.assertContains(response, "S/ 40,00")
        self.assertContains(response, "S/ 60,00")
        self.assertContains(response, "S/ 90,00")
        listed_enrollment = response.context["enrollments"][0]
        self.assertEqual(str(listed_enrollment.total_enrollment_cost), "100.00")
        self.assertEqual(str(listed_enrollment.total_paid), "40.00")
        self.assertEqual(str(listed_enrollment.total_debt), "90.00")

    def test_enrollment_excel_export_returns_workbook_and_respects_filters(self):
        other_student = Student.objects.create(
            campus=self.campus,
            full_name="Otro estudiante para exportar",
        )
        Enrollment.objects.create(
            campus=self.campus,
            student=other_student,
            course=self.course,
        )
        response = self.client.get(
            reverse("enrollment-excel"),
            {"student": self.student.full_name},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn(
            b'attachment; filename="matriculas_', response["Content-Disposition"].encode()
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook["Matrículas"]
        rows = list(worksheet.iter_rows(values_only=True))
        self.assertIn("ESTUDIANTE DE PRUEBA", [value for row in rows for value in row])
        self.assertNotIn("OTRO ESTUDIANTE PARA EXPORTAR", [value for row in rows for value in row])
        self.assertEqual(rows[3][0], "N.º")
        self.assertNotIn("Estado", rows[3])
        self.assertEqual(rows[4][0], 1)

    def test_enrollment_excel_export_excludes_inactive_enrollments(self):
        inactive_student = Student.objects.create(
            campus=self.campus,
            full_name="Estudiante con matrícula inactiva",
        )
        Enrollment.objects.create(
            campus=self.campus,
            student=inactive_student,
            course=self.course,
            active=False,
        )

        response = self.client.get(reverse("enrollment-excel"))

        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook["Matrículas"]
        exported_values = [value for row in worksheet.iter_rows(values_only=True) for value in row]
        self.assertIn(self.student.full_name, exported_values)
        self.assertNotIn(inactive_student.full_name, exported_values)

    def test_enrollment_excel_export_orders_rows_by_student_name(self):
        first_student = Student.objects.create(
            campus=self.campus,
            full_name="Ana Alumna",
        )
        last_student = Student.objects.create(
            campus=self.campus,
            full_name="Zoe Alumna",
        )
        Enrollment.objects.create(
            campus=self.campus,
            student=first_student,
            course=self.course,
        )
        Enrollment.objects.create(
            campus=self.campus,
            student=last_student,
            course=self.course,
        )

        response = self.client.get(reverse("enrollment-excel"))

        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook["Matrículas"]
        rows = list(worksheet.iter_rows(min_row=5, values_only=True))
        self.assertEqual(
            [row[1] for row in rows],
            ["ANA ALUMNA", "ESTUDIANTE DE PRUEBA", "ZOE ALUMNA"],
        )
        self.assertEqual([row[0] for row in rows], [1, 2, 3])

    def test_enrollment_excel_export_uses_debt_status_row_colors(self):
        Installment.objects.create(
            enrollment=self.enrollment,
            concept="Cuota pendiente",
            amount=100,
            due_date=date(2026, 10, 14),
        )
        no_debt_student = Student.objects.create(
            campus=self.campus,
            full_name="Alumno sin deuda",
        )
        Enrollment.objects.create(
            campus=self.campus,
            student=no_debt_student,
            course=self.course,
        )

        response = self.client.get(reverse("enrollment-excel"))

        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook["Matrículas"]
        rows_by_student = {row[1].value: row for row in worksheet.iter_rows(min_row=5)}
        self.assertTrue(
            rows_by_student[self.student.full_name][0].fill.fgColor.rgb.endswith("FFF3CD")
        )
        self.assertTrue(
            rows_by_student[no_debt_student.full_name][0].fill.fgColor.rgb.endswith("D1E7DD")
        )

    def test_enrollment_row_numbers_follow_oldest_enrollment_date(self):
        self.enrollment.enrollment_date = date(2026, 9, 15)
        self.enrollment.save(update_fields=["enrollment_date"])
        older_student = Student.objects.create(
            campus=self.campus,
            full_name="Inscripción más antigua",
        )
        older_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=older_student,
            course=self.course,
            enrollment_date=date(2026, 8, 1),
        )

        response = self.client.get(
            reverse("enrollment-list"),
            {"sort": "student", "direction": "desc"},
        )
        self.assertContains(response, "N.º")
        self.assertContains(response, "?sort=row_number&amp;direction=asc")
        ranks = {
            enrollment.pk: enrollment.row_number for enrollment in response.context["enrollments"]
        }
        self.assertEqual(ranks[older_enrollment.pk], 1)
        self.assertEqual(ranks[self.enrollment.pk], 2)

        ascending = self.client.get(
            reverse("enrollment-list"),
            {"sort": "row_number", "direction": "asc"},
        )
        descending = self.client.get(
            reverse("enrollment-list"),
            {"sort": "row_number", "direction": "desc"},
        )
        self.assertEqual(ascending.context["enrollments"][0].pk, older_enrollment.pk)
        self.assertEqual(descending.context["enrollments"][0].pk, self.enrollment.pk)

        response = self.client.get(reverse("enrollment-excel"))
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook["Matrículas"]
        ranks_by_student = {
            row[1]: row[0] for row in worksheet.iter_rows(min_row=5, values_only=True)
        }
        self.assertEqual(ranks_by_student[self.student.full_name], 1)
        self.assertEqual(ranks_by_student["INSCRIPCIÓN MÁS ANTIGUA"], 2)

    def test_enrollment_filters_cover_student_course_record_code_and_promoter(self):
        promoter = Promoter.objects.create(name="Promotor de prueba", active=True)
        Ficha.objects.create(
            enrollment=self.enrollment,
            campus=self.campus,
            date=self.enrollment.enrollment_date,
            code="42",
        )
        self.enrollment.promoter.add(promoter)

        filters = {
            "student": "Estudiante de",
            "course": "Curso de",
            "record_code": "4",
            "promoter": "Promotor de",
        }
        for field, value in filters.items():
            with self.subTest(field=field):
                response = self.client.get(reverse("enrollment-list"), {field: value})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(len(response.context["enrollments"]), 1)
                self.assertEqual(response.context["enrollments"][0].pk, self.enrollment.pk)

    def test_record_code_check_lists_missing_codes_for_the_selected_date(self):
        check_date = date(2026, 10, 14)
        for record_code in ["1", "2", "4", "5"]:
            student = Student.objects.create(
                campus=self.campus,
                full_name=f"Alumno de ficha {record_code}",
            )
            enrollment = Enrollment.objects.create(
                campus=self.campus,
                student=student,
                course=self.course,
                enrollment_date=check_date,
            )
            Ficha.objects.create(
                enrollment=enrollment, campus=self.campus, date=check_date, code=record_code
            )
        other_date_student = Student.objects.create(
            campus=self.campus,
            full_name="Alumno de otra fecha",
        )
        other_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=other_date_student,
            course=self.course,
            enrollment_date=date(2026, 10, 15),
        )
        Ficha.objects.create(
            enrollment=other_enrollment, campus=self.campus, date=date(2026, 10, 15), code="3"
        )

        response = self.client.get(
            reverse("tools"),
            {
                "record_code_check_date": check_date.isoformat(),
                "record_code_start": "1",
                "record_code_end": "10",
                "check_record_codes": "1",
            },
        )

        self.assertEqual(
            response.context["missing_record_codes"],
            [3, 6, 7, 8, 9, 10],
        )
        self.assertEqual(response.context["missing_record_codes"], [3, 6, 7, 8, 9, 10])
        self.assertContains(response, '<span class="tool-chip">10</span>')

    def test_enrollment_rows_highlight_debt_and_no_debt(self):
        Installment.objects.create(
            enrollment=self.enrollment,
            concept="Pendiente",
            amount=100,
            due_date=date(2026, 10, 14),
        )
        other_student = Student.objects.create(campus=self.campus, full_name="Otro estudiante")
        paid_course = Course.objects.create(
            campus=self.campus,
            name="Curso pagado",
            start_date=date(2026, 10, 14),
        )
        paid_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=other_student,
            course=paid_course,
            active=True,
        )
        paid_installment = Installment.objects.create(
            enrollment=paid_enrollment,
            concept="Pagado",
            amount=50,
            due_date=date(2026, 10, 14),
        )
        paid_payment = Payment.objects.create(
            campus=self.campus,
            enrollment=paid_enrollment,
            cashier=self.user,
            receipt_number="COLOR-001",
            amount=50,
        )
        PaymentAllocation.objects.create(
            payment=paid_payment,
            installment=paid_installment,
            amount=50,
        )

        response = self.client.get(reverse("enrollment-list"))

        self.assertContains(response, "debt-owed")
        self.assertContains(response, "debt-clear")
        self.assertNotContains(response, "row-warning")
        self.assertNotContains(response, "row-success")
        self.assertNotContains(response, "<th>Sede</th>")

    def test_enrollment_list_filters_by_debt_status(self):
        Installment.objects.create(
            enrollment=self.enrollment,
            concept="Pendiente",
            amount=100,
            due_date=date(2026, 10, 14),
        )
        no_debt_student = Student.objects.create(
            campus=self.campus,
            full_name="Estudiante sin deuda",
        )
        no_debt_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=no_debt_student,
            course=self.course,
        )

        with_debt = self.client.get(reverse("enrollment-list"), {"debt_status": "with_debt"})
        without_debt = self.client.get(reverse("enrollment-list"), {"debt_status": "without_debt"})

        self.assertEqual(
            {enrollment.pk for enrollment in with_debt.context["enrollments"]},
            {self.enrollment.pk},
        )
        self.assertEqual(
            {enrollment.pk for enrollment in without_debt.context["enrollments"]},
            {no_debt_enrollment.pk},
        )
        self.assertContains(with_debt, 'value="with_debt" selected')
        self.assertContains(without_debt, 'value="without_debt" selected')

    def test_enrollment_list_filters_by_date_range(self):
        older_student = Student.objects.create(campus=self.campus, full_name="Matrícula anterior")
        older_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=older_student,
            course=self.course,
            enrollment_date=date(2026, 10, 1),
        )

        response = self.client.get(
            reverse("enrollment-list"),
            {"enrollment_date_from": "2026-10-10"},
        )
        self.assertEqual(
            {enrollment.pk for enrollment in response.context["enrollments"]},
            {self.enrollment.pk},
        )
        self.assertContains(response, 'for="enrollment-date-from">Desde')
        self.assertContains(response, 'for="enrollment-date-to">Hasta')

        response = self.client.get(
            reverse("enrollment-list"),
            {"enrollment_date_to": "2026-10-10"},
        )
        self.assertEqual(
            {enrollment.pk for enrollment in response.context["enrollments"]},
            {older_enrollment.pk},
        )

    def test_enrollment_table_sorts_by_student_and_outstanding_debt(self):
        other_student = Student.objects.create(campus=self.campus, full_name="Zeta Estudiante")
        other_enrollment = Enrollment.objects.create(
            campus=self.campus,
            student=other_student,
            course=self.course,
        )

        response = self.client.get(
            reverse("enrollment-list"),
            {"sort": "student", "direction": "desc"},
        )
        self.assertEqual(response.context["enrollments"][0].pk, other_enrollment.pk)
        self.assertContains(response, "?sort=student&amp;direction=asc")

        self.add_payment_history()
        Installment.objects.create(
            enrollment=other_enrollment,
            concept="Deuda menor",
            amount=30,
            due_date=date(2026, 10, 14),
        )
        response = self.client.get(
            reverse("enrollment-list"),
            {"sort": "total_debt", "direction": "desc"},
        )
        self.assertEqual(response.context["enrollments"][0].pk, self.enrollment.pk)

    def test_enrollment_filters_are_preserved_in_pagination(self):
        for index in range(26):
            student = Student.objects.create(
                campus=self.campus,
                full_name=f"Alumno paginado {index}",
            )
            enrollment = Enrollment.objects.create(
                campus=self.campus,
                student=student,
                course=self.course,
            )
            Ficha.objects.create(
                enrollment=enrollment,
                campus=self.campus,
                date=enrollment.enrollment_date,
                code=f"999{index}",
            )

        response = self.client.get(reverse("enrollment-list"), {"record_code": "999"})

        self.assertTrue(response.context["is_paginated"])
        self.assertContains(response, "?record_code=999&amp;page=2")

    def test_enrollment_list_shows_associated_receipt_numbers_and_void_status(self):
        confirmed_payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="RECIBO-ACTIVO-01",
            amount=20,
        )
        void_payment = Payment.objects.create(
            campus=self.campus,
            enrollment=self.enrollment,
            cashier=self.user,
            receipt_number="RECIBO-ANULADO-02",
            amount=15,
            status=Payment.Status.VOID,
        )

        response = self.client.get(reverse("enrollment-list"))

        self.assertContains(response, "Recibos")
        self.assertContains(response, "RECIBO-ACTIVO-01")
        self.assertContains(response, "RECIBO-ANULADO-02")
        self.assertContains(response, "Anulado")
        self.assertContains(
            response,
            f'href="{reverse("payment-receipt", args=[confirmed_payment.pk])}"',
        )
        self.assertContains(
            response,
            f'href="{reverse("payment-receipt", args=[void_payment.pk])}"',
        )
        self.assertEqual(
            self.client.get(reverse("payment-receipt", args=[void_payment.pk])).status_code,
            200,
        )

    def test_student_form_accepts_multiple_phones_and_labels_document_as_dni(self):
        response = self.client.get(reverse("student-update", args=[self.student.pk]))
        self.assertContains(response, "DNI")
        self.assertContains(response, "Teléfonos", count=1)
        response = self.client.post(
            reverse("student-update", args=[self.student.pk]),
            {
                "last_name": "Estudiante de",
                "first_name": "prueba",
                "document_number": self.student.document_number,
                "active": "on",
                **self.phone_formset_data(["987654321", "912345678"]),
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.student.refresh_from_db()
        self.assertEqual(self.student.phone, "987654321, 912345678")
        self.assertEqual(
            list(self.student.phones.values_list("number", flat=True)),
            ["987654321", "912345678"],
        )

    def test_student_form_requires_at_least_one_phone(self):
        response = self.client.post(
            reverse("student-update", args=[self.student.pk]),
            {
                "last_name": "Estudiante de",
                "first_name": "prueba",
                "document_number": self.student.document_number,
                "active": "on",
                **self.phone_formset_data([""]),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Este campo es obligatorio.")
        self.student.refresh_from_db()
        self.assertEqual(self.student.phone, "987654321")

    def test_student_form_rejects_any_invalid_phone(self):
        response = self.client.post(
            reverse("student-update", args=[self.student.pk]),
            {
                "last_name": "Estudiante de",
                "first_name": "prueba",
                "document_number": self.student.document_number,
                "active": "on",
                **self.phone_formset_data(["987654321", "123456789"]),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "celular peruano")
        self.student.refresh_from_db()
        self.assertEqual(self.student.phone, "987654321")

    def test_student_form_can_remove_a_phone_individually(self):
        data = self.phone_formset_data(["987654321", "912345678"])
        data["phones-0-DELETE"] = "on"
        response = self.client.post(
            reverse("student-update", args=[self.student.pk]),
            {
                "last_name": "Estudiante de",
                "first_name": "prueba",
                "document_number": self.student.document_number,
                "active": "on",
                **data,
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.student.refresh_from_db()
        self.assertEqual(self.student.phone, "912345678")
        self.assertEqual(list(self.student.phones.values_list("number", flat=True)), ["912345678"])

    def test_legacy_comma_separated_phones_convert_to_related_rows(self):
        student = Student.objects.create(
            campus=self.campus,
            full_name="Estudiante importado",
            phone="987654321, 912345678, 987654321",
        )

        create_student_phones_from_legacy(student)

        self.assertEqual(
            list(student.phones.values_list("number", flat=True)),
            ["987654321", "912345678"],
        )
        student.refresh_from_db()
        self.assertEqual(student.phone, "987654321, 912345678")


class PendingEnrollmentTests(TestCase):
    def test_course_and_enrollment_fees_must_be_greater_than_zero(self):
        campus = Campus.objects.create(name="Sede Montos")
        course = Course(campus=campus, name="X", schedule="Y", enrollment_fee=0, monthly_fee=-5)
        with self.assertRaises(ValidationError) as context:
            course.full_clean()
        self.assertIn("enrollment_fee", context.exception.message_dict)
        self.assertIn("monthly_fee", context.exception.message_dict)
        enrollment = Enrollment(enrollment_fee=0, monthly_fee=0)
        with self.assertRaises(ValidationError) as context:
            enrollment.full_clean()
        self.assertIn("enrollment_fee", context.exception.message_dict)
        self.assertIn("monthly_fee", context.exception.message_dict)

    def test_new_campus_gets_default_promoters(self):
        campus = Campus.objects.create(name="Sede Promotores")
        names = set(campus.promoters.values_list("name", flat=True))
        self.assertEqual(names, {"PARTICULAR", "NN"})
        other = Campus.objects.create(name="Otra Sede Promotores")
        self.assertEqual(Promoter.objects.filter(name="NN").count(), 1)
        self.assertTrue(other.promoters.filter(name="NN").exists())

    def test_new_campus_does_not_get_a_placeholder_course_or_schedule(self):
        campus = Campus.objects.create(name="Sede Nueva")
        self.assertFalse(Course.objects.filter(campus=campus).exists())
        self.assertFalse(Schedule.objects.filter(campus=campus, name="-").exists())
        self.assertEqual(
            set(Schedule.objects.filter(campus=campus).values_list("name", flat=True)),
            {"SÁBADOS DE 8AM A 11AM", "SÁBADOS DE 11AM A 2PM"},
        )

    def test_enrollment_must_have_a_student_or_a_course(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Enrollment.objects.create(
                    campus=Campus.objects.create(name="Sede sin asignación"),
                    student=None,
                    course=None,
                )


class CourseUniquenessTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede de cursos", active=True)
        self.course_data = {
            "campus": self.campus,
            "name": "Matemáticas",
            "schedule": "Lunes 18:00",
            "group": "A",
        }

    def test_course_group_must_be_positive_integer(self):
        for bad in ("A", "0", "-1", "1.5", ""):
            with self.assertRaises(ValidationError, msg=bad) as context:
                Course(**{**self.course_data, "group": bad}).full_clean()
            self.assertIn("group", context.exception.message_dict)
        fees = {"enrollment_fee": 10, "monthly_fee": 10}
        Course(**{**self.course_data, "group": "2", **fees}).full_clean()

    def test_same_course_identity_cannot_be_created_twice_in_one_campus(self):
        Course.objects.create(**self.course_data)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Course.objects.create(**self.course_data)

    def test_same_course_identity_is_allowed_in_another_campus(self):
        Course.objects.create(**self.course_data)
        other_campus = Campus.objects.create(name="Otra sede", active=True)
        Course.objects.create(**{**self.course_data, "campus": other_campus})

    def test_course_requires_schedule_and_defaults_group_to_one(self):
        course = Course(campus=self.campus, name="Curso sin horario")

        with self.assertRaises(ValidationError) as context:
            course.full_clean()

        self.assertIn("schedule", context.exception.message_dict)
        self.assertEqual(course.group, "1")

    def test_default_schedules_are_available_in_the_catalog(self):
        self.assertEqual(
            set(Schedule.objects.filter(campus=self.campus).values_list("name", flat=True)),
            {"SÁBADOS DE 8AM A 11AM", "SÁBADOS DE 11AM A 2PM"},
        )


class StudentMergeTests(TestCase):
    def setUp(self):
        from calls.models import CallNote

        self.CallNote = CallNote
        self.campus = Campus.objects.create(name="Sede Fusión", active=True)
        self.user = create_user(username="merger", password="x" * 12, manage_students=True)
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.keep = Student.objects.create(campus=self.campus, full_name="Juan Pérez", phone="999")
        self.dup = Student.objects.create(campus=self.campus, full_name="Perez Juan", phone="111")
        Student.objects.create(campus=self.campus, full_name="Otra Persona", phone="222")
        self.course = Course.objects.create(campus=self.campus, name="Curso M")
        self.enrollment = Enrollment.objects.create(
            campus=self.campus, student=self.dup, course=self.course
        )
        self.note = CallNote.objects.create(
            campus=self.campus, student=self.dup, author=self.user, note="hola"
        )

    def merge(self, keep, duplicate):
        return self.client.post(
            reverse("student-duplicates"), {"keep": keep.pk, "duplicate": duplicate.pk}
        )

    def test_duplicates_are_grouped_ignoring_accents_and_word_order(self):
        groups = self.client.get(reverse("student-duplicates")).context["groups"]
        self.assertEqual([{s.pk for s in g} for g in groups], [{self.keep.pk, self.dup.pk}])

    def test_merge_moves_enrollments_and_calls_and_deactivates_duplicate(self):
        self.merge(self.keep, self.dup)
        self.enrollment.refresh_from_db()
        self.note.refresh_from_db()
        self.dup.refresh_from_db()
        self.assertEqual(self.enrollment.student, self.keep)
        self.assertEqual(self.note.student, self.keep)
        self.assertFalse(self.dup.active)
        self.assertEqual(self.client.get(reverse("student-duplicates")).context["groups"], [])

    def test_merge_preserves_distinct_phone_records_and_updates_compatibility_value(self):
        StudentPhone.objects.create(student=self.keep, number="988888888")
        StudentPhone.objects.create(student=self.dup, number="977777777")
        StudentPhone.objects.create(student=self.dup, number="988888888")

        self.merge(self.keep, self.dup)

        self.keep.refresh_from_db()
        self.assertEqual(
            list(self.keep.phones.values_list("number", flat=True)),
            ["988888888", "977777777"],
        )
        self.assertEqual(self.keep.phone, "988888888, 977777777")
        self.assertFalse(self.dup.phones.exists())

    def test_merge_rejects_different_people_and_requires_permission(self):
        other = Student.objects.get(full_name="OTRA PERSONA")
        self.merge(self.keep, other)
        other.refresh_from_db()
        self.assertTrue(other.active)
        self.user.groups.clear()
        self.user._perm_cache = None
        self.assertEqual(self.client.get(reverse("student-duplicates")).status_code, 403)


class StudentImportAndTimelineTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede Import", active=True)
        self.user = create_user(
            username="importer", manage_students=True, register_payments=True, manage_calls=True
        )
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        Student.objects.create(campus=self.campus, full_name="Juan Pérez", phone="911111111")
        self.course = Course.objects.create(
            campus=self.campus, name="Inglés", enrollment_fee=30, monthly_fee=50
        )

    def upload(self, rows):
        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["Nombre", "DNI", "Teléfono", "Curso"])
        for row in rows:
            book.active.append(row)
        buffer = BytesIO()
        book.save(buffer)
        buffer.name = "alumnos.xlsx"
        buffer.seek(0)
        return self.client.post(reverse("student-import"), {"archivo": buffer})

    def test_preview_classifies_rows_and_confirm_creates_students_and_enrollments(self):
        response = self.upload(
            [
                ["Ana Torres", "12345678", "922222222", "Inglés"],
                ["Perez Juan", "", "", ""],
                ["Luis Ruiz", "", "123", ""],
                ["Eva Soto", "", "", "No existe"],
            ]
        )
        self.assertEqual(response.context["counts"], {"new": 1, "duplicate": 1, "error": 2})
        self.assertEqual(Student.objects.filter(campus=self.campus).count(), 1)

        self.client.post(reverse("student-import"), {"accion": "confirmar"})

        ana = Student.objects.get(full_name="ANA TORRES")
        enrollment = Enrollment.objects.get(student=ana)
        self.assertEqual(enrollment.course, self.course)
        self.assertEqual(list(ana.phones.values_list("number", flat=True)), ["922222222"])
        self.assertEqual(ana.phone, "922222222")
        self.assertEqual(Installment.objects.filter(enrollment=enrollment).count(), 2)
        self.assertEqual(Student.objects.filter(campus=self.campus).count(), 2)

    def test_import_requires_permission_and_rejects_non_excel(self):
        revoke(self.user, "manage_students")
        self.assertEqual(self.client.get(reverse("student-import")).status_code, 403)

    def test_timeline_merges_enrollments_payments_and_calls(self):
        from calls.models import CallNote

        student = Student.objects.get(full_name="JUAN PÉREZ")
        enrollment = Enrollment.objects.create(
            campus=self.campus, student=student, course=self.course
        )
        Payment.objects.create(
            campus=self.campus,
            enrollment=enrollment,
            cashier=self.user,
            receipt_number="T1",
            amount=10,
        )
        CallNote.objects.create(campus=self.campus, student=student, author=self.user, note="hola")

        response = self.client.get(reverse("student-timeline", args=[student.pk]))

        kinds = {event["kind"] for event in response.context["events"]}
        self.assertTrue({"Matrícula", "Pago", "Llamada"} <= kinds)
        revoke(self.user, "register_payments")
        kinds = {
            e["kind"]
            for e in self.client.get(reverse("student-timeline", args=[student.pk])).context[
                "events"
            ]
        }
        self.assertNotIn("Pago", kinds)

    def test_profile_shows_full_record_with_payments_and_history(self):
        from calls.models import CallNote
        from core.models import ActivityLog

        student = Student.objects.get(full_name="JUAN PÉREZ")
        StudentPhone.objects.create(student=student, number="911111111")
        current = Enrollment.objects.create(campus=self.campus, student=student, course=self.course)
        old_course = Course.objects.create(
            campus=self.campus, name="Inglés Básico", schedule="L", group="2"
        )
        old = Enrollment.objects.create(
            campus=self.campus, student=student, course=old_course, active=False
        )
        installment = Installment.objects.create(
            enrollment=current, concept="Cuota 1", amount=100, due_date=date(2030, 1, 1)
        )
        payment = Payment.objects.create(
            campus=self.campus,
            enrollment=current,
            cashier=self.user,
            receipt_number="P1",
            amount=40,
        )
        PaymentAllocation.objects.create(payment=payment, installment=installment, amount=40)
        CallNote.objects.create(
            campus=self.campus, student=student, author=self.user, note="Llamar mañana"
        )
        ActivityLog.objects.create(
            user=self.user,
            username=self.user.username,
            action="update",
            description="Editó un estudiante",
            object_type="student",
            object_id=student.pk,
        )

        response = self.client.get(reverse("student-profile", args=[student.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual([c["enrollment"] for c in response.context["current"]], [current])
        self.assertEqual([c["enrollment"] for c in response.context["previous"]], [old])
        card = response.context["current"][0]
        self.assertEqual((card["total"], card["paid"], card["balance"]), (100, 40, 60))
        self.assertContains(response, "911111111")
        self.assertContains(response, "P1")
        self.assertContains(response, "Llamar mañana")
        who = {e["who"] for e in response.context["events"]}
        self.assertIn(self.user.username, who)

    def test_profile_hides_sections_without_permission(self):
        student = Student.objects.get(full_name="JUAN PÉREZ")
        enrollment = Enrollment.objects.create(
            campus=self.campus, student=student, course=self.course
        )
        Payment.objects.create(
            campus=self.campus,
            enrollment=enrollment,
            cashier=self.user,
            receipt_number="SECRETO1",
            amount=10,
        )
        revoke(self.user, "register_payments")

        response = self.client.get(reverse("student-profile", args=[student.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "SECRETO1")
        self.assertFalse(response.context["can_payments"])

    def test_profile_is_scoped_to_campus_and_requires_permission(self):
        other_campus = Campus.objects.create(name="Otra sede", active=True)
        foreign = Student.objects.create(campus=other_campus, full_name="AJENO", phone="922222222")
        self.assertEqual(
            self.client.get(reverse("student-profile", args=[foreign.pk])).status_code, 404
        )
        student = Student.objects.get(full_name="JUAN PÉREZ")
        revoke(self.user, "manage_students")
        self.assertEqual(
            self.client.get(reverse("student-profile", args=[student.pk])).status_code, 403
        )


class StudentImportErrorPathTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede Errores", active=True)
        self.user = create_user(username="imp2")
        grant(self.user, "students.add_student")
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()

    def post_file(self, name, content):
        upload = BytesIO(content)
        upload.name = name
        return self.client.post(reverse("student-import"), {"archivo": upload}, follow=True)

    def xlsx(self, rows):
        from openpyxl import Workbook

        book = Workbook()
        for row in rows:
            book.active.append(row)
        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    def test_rejects_missing_or_non_excel_files(self):
        response = self.client.post(reverse("student-import"), follow=True)
        self.assertContains(response, "Sube un archivo Excel")
        response = self.post_file("alumnos.csv", b"Nombre\nAna")
        self.assertContains(response, "Sube un archivo Excel")

    def test_rejects_corrupt_excel_and_missing_name_column(self):
        response = self.post_file("roto.xlsx", b"no es un excel")
        self.assertContains(response, "No se pudo leer el archivo")
        response = self.post_file("sin_nombre.xlsx", self.xlsx([["DNI"], ["123"]]))
        self.assertContains(response, "Falta la columna")
        self.assertEqual(Student.objects.count(), 0)

    def test_confirm_without_preview_creates_nothing(self):
        response = self.client.post(reverse("student-import"), {"accion": "confirmar"}, follow=True)
        self.assertContains(response, "No hay una vista previa pendiente")
        self.assertEqual(Student.objects.count(), 0)

    def test_enrollments_need_permission_and_repeated_names_in_file_are_duplicates(self):
        course = Course.objects.create(campus=self.campus, name="Arte", monthly_fee=10)
        response = self.post_file(
            "alumnos.xlsx",
            self.xlsx(
                [
                    ["Nombre", "Telefono", "Curso"],
                    ["Ana Gil", "955555555", "Arte"],
                    ["Gil Ana", "944444444", ""],
                    [None, None, None],
                ]
            ),
        )
        self.assertEqual(response.context["counts"], {"new": 1, "duplicate": 1, "error": 0})
        self.client.post(reverse("student-import"), {"accion": "confirmar"})
        self.assertEqual(Student.objects.count(), 1)
        self.assertEqual(Enrollment.objects.filter(course=course).count(), 0)


class StudentSearchTests(TestCase):
    def setUp(self):
        self.campus = Campus.objects.create(name="Sede Búsqueda", active=True)
        self.user = create_user(username="buscador", manage_students=True)
        self.user.campuses.add(self.campus)
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_campus_id"] = self.campus.pk
        session.save()
        self.ana = Student.objects.create(
            campus=self.campus,
            full_name="María Pérez Soto",
            document_number="45678912",
            phone="987654321",
        )
        Student.objects.create(campus=self.campus, full_name="Luis Ramos", phone="911222333")

    def names(self, q):
        response = self.client.get(reverse("student-list"), {"q": q})
        return [s.full_name for s in response.context["students"]]

    def test_matches_name_in_any_order_without_accents_dni_and_phone(self):
        self.assertEqual(self.names("perez maria"), ["MARÍA PÉREZ SOTO"])
        self.assertEqual(self.names("soto"), ["MARÍA PÉREZ SOTO"])
        self.assertEqual(self.names("456789"), ["MARÍA PÉREZ SOTO"])
        self.assertEqual(self.names("98765"), ["MARÍA PÉREZ SOTO"])
        self.assertEqual(self.names("maria ramos"), [])

    def test_matches_ficha_code(self):
        course = Course.objects.create(campus=self.campus, name="Arte")
        enrollment = Enrollment.objects.create(campus=self.campus, student=self.ana, course=course)
        Ficha.objects.create(
            enrollment=enrollment, campus=self.campus, date=enrollment.enrollment_date, code=4321
        )
        self.assertEqual(self.names("4321"), ["MARÍA PÉREZ SOTO"])
        self.assertEqual(self.names("9999"), [])


class EnrollmentUniquenessTests(TestCase):
    def test_database_rejects_duplicate_student_course(self):
        campus = Campus.objects.create(name="Sede unica", active=True)
        student = Student.objects.create(campus=campus, full_name="Alumno Unico")
        course = Course.objects.create(
            campus=campus, name="Curso unico", start_date=date(2026, 10, 14)
        )
        Enrollment.objects.create(campus=campus, student=student, course=course)

        with self.assertRaises(ValidationError):
            Enrollment(campus=campus, student=student, course=course).full_clean()
        with self.assertRaises(IntegrityError), transaction.atomic():
            Enrollment.objects.create(campus=campus, student=student, course=course)

    def test_pending_enrollments_are_not_restricted(self):
        campus = Campus.objects.create(name="Sede pendientes", active=True)
        student = Student.objects.create(campus=campus, full_name="Alumno Pendiente")
        Enrollment.objects.create(campus=campus, student=student)
        Enrollment.objects.create(campus=campus, student=student)


@skipUnlessDBFeature("has_select_for_update")
class ConcurrentEnrollmentTests(TransactionTestCase):
    reset_sequences = True

    def test_concurrent_duplicate_enrollments_are_rejected_by_the_database(self):
        campus = Campus.objects.create(name="Sede matrícula concurrente")
        student = Student.objects.create(campus=campus, full_name="Alumno único concurrente")
        course = Course.objects.create(
            campus=campus, name="Curso único concurrente", start_date=date(2026, 10, 14)
        )
        barrier = threading.Barrier(2)
        results = []
        errors = []

        def create_enrollment():
            connections.close_all()
            try:
                barrier.wait(timeout=10)
                with transaction.atomic():
                    Enrollment.objects.create(campus=campus, student=student, course=course)
                results.append("created")
            except IntegrityError:
                results.append("duplicate")
            except Exception as error:
                errors.append(error)
            finally:
                connections.close_all()

        threads = [threading.Thread(target=create_enrollment) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)

        self.assertFalse(any(thread.is_alive() for thread in threads), "Enrollment request hung.")
        self.assertEqual(errors, [])
        self.assertCountEqual(results, ["created", "duplicate"])
        self.assertEqual(
            Enrollment.objects.filter(student=student, course=course).count(),
            1,
        )
