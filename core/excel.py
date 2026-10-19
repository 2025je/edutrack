from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = '"S/" #,##0.00'
DATE_FORMAT = "DD/MM/YYYY"


def xlsx_response(filename, title, subtitle, columns, rows, *, money=(), dates=(), totals=None):
    """Arma un Excel con título, tabla filtrable y formatos de moneda y fecha.

    `money` y `dates` son los índices (desde 0) de las columnas a formatear.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(columns))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    sheet["A1"].fill = PatternFill("solid", fgColor="173A72")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = f"{subtitle} · Exportado: {timezone.localdate():%d/%m/%Y}"

    for index, header in enumerate(columns, start=1):
        cell = sheet.cell(row=4, column=index, value=header)
        cell.font = Font(bold=True, color="173A72")
        cell.fill = PatternFill("solid", fgColor="EAF1FF")
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(len(header) + 6, 40))

    count = 0
    for row in rows:
        sheet.append(list(row))
        count += 1
    for column in money:
        for cell in sheet.iter_cols(min_col=column + 1, max_col=column + 1, min_row=5).__next__():
            cell.number_format = MONEY_FORMAT
    for column in dates:
        for cell in sheet.iter_cols(min_col=column + 1, max_col=column + 1, min_row=5).__next__():
            cell.number_format = DATE_FORMAT

    if count:
        table = Table(displayName="Datos", ref=f"A4:{last_column}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        if totals:
            sheet.append(list(totals))
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
                if cell.column - 1 in money:
                    cell.number_format = MONEY_FORMAT
    else:
        sheet["A4"] = "No hay datos para exportar con estos criterios."
    sheet.freeze_panes = "A5"

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type=XLSX_TYPE)
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}_{timezone.localdate():%Y%m%d}.xlsx"'
    )
    return response
